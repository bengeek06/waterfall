"""Parse an Excel reconciliation workbook back into typed rows (E6-09, #70).

Mirror of ``estimate_reconciliation_export.py``: reads the exact three-sheet
layout (``Tâches`` / ``MO`` / ``Non-MO``) that module writes, sharing its
column headers (``TASK_HEADERS`` / ``LABOR_HEADERS`` / ``COST_LINE_HEADERS``)
as the single source of truth for both directions of the round trip.

Deliberately narrow scope: this module only turns workbook bytes into typed,
per-row dataclasses (or raises :class:`EstimateReconciliationFormatError` when
the file cannot even be read as structured data). It has no FastAPI
dependency and never touches the database -- the actual reconciliation
diff/apply logic (which needs project/estimate context, existing rows, and
route-layer helpers like ``resolve_cost_code_id``) lives in
``api/routes/estimates.py``, next to the two routes that call it
(``.../import-reconciliation/preview`` and ``.../import-reconciliation/confirm``).

Two classes of problem are told apart on purpose:

* A **format** problem (missing sheet, header row that doesn't match the
  exported layout, or a cell holding a value of the wrong type where a
  specific one is required, e.g. text in a numeric column) means the file
  itself cannot be trusted as structured data -- collected into
  :class:`EstimateReconciliationFormatError` (mirroring
  ``waterfall.services.msproject_xml.MsProjectValidationError``'s shape) and
  reported as a single 400 with the full list of problems, never a partial
  parse.
* A **business** problem (a blank cell where a value is only required for
  *some* rows, e.g. ``label`` on a new Non-MO line but not on one already
  flagged for deletion) is left as ``None`` here and turned into a precisely
  attributed :class:`~waterfall.schemas.projects.ReconciliationIssue` by the
  diff/apply step instead, since only that step knows whether the row is a
  create, an update, or a mutation-ignored existing row.

The revision half (E14-07c, #365)
---------------------------------
Below the legacy parser, this module also reads the revision workbook
(:func:`parse_revision_reconciliation_workbook`) **and** reconciles it against a
revision (:func:`reconcile_revision`). That second half is deliberately not a pure
parser, and the reason the legacy one is stays instructive: the legacy diff/apply
had to live in ``api/routes/estimates.py`` because it needed route-layer helpers
and three tables' worth of staging. The revision one needs neither -- every rule
it enforces is a *domain* rule, called through
:mod:`waterfall.services.revision_store`, and the only two things it does that the
domain cannot are reading the referential (is this an active labour role) and
deciding which nodes the file no longer mentions. That is service work, and it
belongs next to the parser whose output it consumes rather than in a handler.

Nothing in the revision half commits: the caller owns the transaction, as in
:mod:`waterfall.services.revision_import`.
"""

from __future__ import annotations

from dataclasses import dataclass, replace
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from waterfall.domain import revision as domain
from waterfall.domain.revision.pricing import AmountResolver
from waterfall.models.resources import CostCategory, CostType, ProjectCostCode, ResourceRole
from waterfall.schemas.resources import CostTypeKind
from waterfall.services.estimate_calculation import price_loaded_revision
from waterfall.services.estimate_reconciliation_export import (
    COST_LINE_HEADERS,
    LABOR_HEADERS,
    REVISION_COST_LINE_HEADERS,
    REVISION_LABOR_HEADERS,
    REVISION_TASK_HEADERS,
    TASK_HEADERS,
)
from waterfall.services.revision_store import LoadedRevision, load_revision, save_revision
from waterfall.services.revision_tree import RevisionLockConflictError

_TASK_SHEET = "Tâches"
_LABOR_SHEET = "MO"
_COST_LINE_SHEET = "Non-MO"


class EstimateReconciliationFormatError(ValueError):
    """The uploaded workbook cannot be parsed into reconciliation rows.

    ``issues`` is a list of structured dicts (``code``/``message``, plus
    ``sheet`` and, for a row-scoped problem, ``row``) -- the same shape as
    ``waterfall.services.msproject_xml.MsProjectValidationError.issues`` --
    collected across the *whole* workbook before raising, so a caller gets
    every format problem in one round trip instead of fixing them one at a
    time.
    """

    def __init__(self, issues: list[dict[str, object]]) -> None:
        self.issues = issues
        super().__init__("Estimate reconciliation workbook is invalid")


@dataclass(frozen=True)
class TaskFileRow:
    row_number: int
    id: int | None
    task_id: int | None
    parent_task_id: int | None
    position: int | None
    task_name: str | None
    outline_number: str | None
    outline_level: int | None
    is_milestone: bool


@dataclass(frozen=True)
class LaborFileRow:
    row_number: int
    id: int | None
    task_id: int | None
    role_id: int | None
    cost_code_id: int | None
    quantity: Decimal | None
    hours: Decimal | None
    comment: str | None
    hors_perimetre_planning: bool


@dataclass(frozen=True)
class CostLineFileRow:
    row_number: int
    id: int | None
    task_id: int | None
    cost_category_id: int | None
    cost_code_id: int | None
    label: str | None
    quantity: Decimal | None
    unit_cost: Decimal | None
    supply_status: str | None
    planned_date: datetime | None


@dataclass(frozen=True)
class ParsedReconciliationWorkbook:
    tasks: list[TaskFileRow]
    labor: list[LaborFileRow]
    cost_lines: list[CostLineFileRow]


def _is_blank_row(values: tuple[object, ...]) -> bool:
    return all(value is None or value == "" for value in values)


def _pad(values: tuple[object, ...], length: int) -> tuple[object, ...]:
    if len(values) >= length:
        return values
    return values + (None,) * (length - len(values))


def _cell_int(
    value: object,
    *,
    sheet: str,
    row_number: int,
    column: str,
    issues: list[dict[str, object]],
) -> int | None:
    if value is None or value == "":
        return None
    if not isinstance(value, bool):
        if isinstance(value, int):
            return value
        if isinstance(value, float) and value.is_integer():
            return int(value)
        if isinstance(value, str):
            stripped = value.strip()
            if stripped.lstrip("-").isdigit():
                return int(stripped)
    issues.append(
        {
            "code": "INVALID_INTEGER",
            "message": f"Column '{column}' must be a whole number",
            "sheet": sheet,
            "row": row_number,
        }
    )
    return None


def _cell_decimal(
    value: object,
    *,
    sheet: str,
    row_number: int,
    column: str,
    issues: list[dict[str, object]],
) -> Decimal | None:
    if value is None or value == "":
        return None
    if not isinstance(value, bool):
        if isinstance(value, int | float):
            return Decimal(str(value))
        if isinstance(value, str):
            try:
                return Decimal(value.strip())
            except InvalidOperation:
                pass
    issues.append(
        {
            "code": "INVALID_NUMBER",
            "message": f"Column '{column}' must be a number",
            "sheet": sheet,
            "row": row_number,
        }
    )
    return None


def _cell_bool(value: object) -> bool:
    if isinstance(value, bool):
        return value
    if value is None or value == "":
        return False
    if isinstance(value, int | float):
        return bool(value)
    if isinstance(value, str):
        return value.strip().lower() in {"true", "1", "vrai", "oui", "yes"}
    return False


def _cell_text(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _cell_datetime(
    value: object,
    *,
    sheet: str,
    row_number: int,
    column: str,
    issues: list[dict[str, object]],
) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day, tzinfo=UTC)
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value.strip())
        except ValueError:
            pass
    issues.append(
        {
            "code": "INVALID_DATE",
            "message": f"Column '{column}' must be an ISO date",
            "sheet": sheet,
            "row": row_number,
        }
    )
    return None


def _check_headers(
    workbook: Workbook,
    sheet_name: str,
    expected: list[str],
    issues: list[dict[str, object]],
) -> None:
    if sheet_name not in workbook.sheetnames:
        issues.append(
            {
                "code": "MISSING_SHEET",
                "message": f"Missing sheet '{sheet_name}'",
                "sheet": sheet_name,
            }
        )
        return
    sheet = workbook[sheet_name]
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    actual = list(_pad(tuple(header_row), len(expected)))[: len(expected)]
    if actual != expected:
        issues.append(
            {
                "code": "INVALID_HEADER",
                "message": f"Sheet '{sheet_name}' must start with columns {expected}",
                "sheet": sheet_name,
            }
        )


def _parse_task_rows(sheet: Worksheet, issues: list[dict[str, object]]) -> list[TaskFileRow]:
    rows: list[TaskFileRow] = []
    for row_number, raw_values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        values = _pad(tuple(raw_values), len(TASK_HEADERS))
        if _is_blank_row(values):
            continue
        (
            raw_id,
            raw_task_id,
            _raw_task_uid,
            raw_parent_task_id,
            raw_position,
            raw_task_name,
            raw_outline_number,
            raw_outline_level,
            raw_is_milestone,
        ) = values[: len(TASK_HEADERS)]
        rows.append(
            TaskFileRow(
                row_number=row_number,
                id=_cell_int(
                    raw_id, sheet=_TASK_SHEET, row_number=row_number, column="id", issues=issues
                ),
                task_id=_cell_int(
                    raw_task_id,
                    sheet=_TASK_SHEET,
                    row_number=row_number,
                    column="task_id",
                    issues=issues,
                ),
                parent_task_id=_cell_int(
                    raw_parent_task_id,
                    sheet=_TASK_SHEET,
                    row_number=row_number,
                    column="parent_task_id",
                    issues=issues,
                ),
                position=_cell_int(
                    raw_position,
                    sheet=_TASK_SHEET,
                    row_number=row_number,
                    column="position",
                    issues=issues,
                ),
                task_name=_cell_text(raw_task_name),
                outline_number=_cell_text(raw_outline_number),
                outline_level=_cell_int(
                    raw_outline_level,
                    sheet=_TASK_SHEET,
                    row_number=row_number,
                    column="outline_level",
                    issues=issues,
                ),
                is_milestone=_cell_bool(raw_is_milestone),
            )
        )
    return rows


def _parse_labor_rows(sheet: Worksheet, issues: list[dict[str, object]]) -> list[LaborFileRow]:
    rows: list[LaborFileRow] = []
    for row_number, raw_values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        values = _pad(tuple(raw_values), len(LABOR_HEADERS))
        if _is_blank_row(values):
            continue
        (
            raw_id,
            raw_task_id,
            _raw_task_name,
            raw_role_id,
            _raw_role_name,
            _raw_cost_category_id,
            _raw_cost_category_name,
            raw_cost_code_id,
            raw_quantity,
            raw_hours,
            raw_comment,
            raw_hors_perimetre,
        ) = values[: len(LABOR_HEADERS)]
        rows.append(
            LaborFileRow(
                row_number=row_number,
                id=_cell_int(
                    raw_id, sheet=_LABOR_SHEET, row_number=row_number, column="id", issues=issues
                ),
                task_id=_cell_int(
                    raw_task_id,
                    sheet=_LABOR_SHEET,
                    row_number=row_number,
                    column="task_id",
                    issues=issues,
                ),
                role_id=_cell_int(
                    raw_role_id,
                    sheet=_LABOR_SHEET,
                    row_number=row_number,
                    column="role_id",
                    issues=issues,
                ),
                cost_code_id=_cell_int(
                    raw_cost_code_id,
                    sheet=_LABOR_SHEET,
                    row_number=row_number,
                    column="cost_code_id",
                    issues=issues,
                ),
                quantity=_cell_decimal(
                    raw_quantity,
                    sheet=_LABOR_SHEET,
                    row_number=row_number,
                    column="quantity",
                    issues=issues,
                ),
                hours=_cell_decimal(
                    raw_hours,
                    sheet=_LABOR_SHEET,
                    row_number=row_number,
                    column="hours",
                    issues=issues,
                ),
                comment=_cell_text(raw_comment),
                hors_perimetre_planning=_cell_bool(raw_hors_perimetre),
            )
        )
    return rows


def _parse_cost_line_rows(
    sheet: Worksheet, issues: list[dict[str, object]]
) -> list[CostLineFileRow]:
    rows: list[CostLineFileRow] = []
    for row_number, raw_values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        values = _pad(tuple(raw_values), len(COST_LINE_HEADERS))
        if _is_blank_row(values):
            continue
        (
            raw_id,
            raw_task_id,
            _raw_cost_type_id,
            _raw_cost_type_code,
            raw_cost_category_id,
            _raw_accounting_code,
            _raw_category_code,
            raw_cost_code_id,
            raw_label,
            raw_quantity,
            raw_unit_cost,
            _raw_purchase_cost,
            raw_supply_status,
            raw_planned_date,
        ) = values[: len(COST_LINE_HEADERS)]
        rows.append(
            CostLineFileRow(
                row_number=row_number,
                id=_cell_int(
                    raw_id,
                    sheet=_COST_LINE_SHEET,
                    row_number=row_number,
                    column="id",
                    issues=issues,
                ),
                task_id=_cell_int(
                    raw_task_id,
                    sheet=_COST_LINE_SHEET,
                    row_number=row_number,
                    column="task_id",
                    issues=issues,
                ),
                cost_category_id=_cell_int(
                    raw_cost_category_id,
                    sheet=_COST_LINE_SHEET,
                    row_number=row_number,
                    column="cost_category_id",
                    issues=issues,
                ),
                cost_code_id=_cell_int(
                    raw_cost_code_id,
                    sheet=_COST_LINE_SHEET,
                    row_number=row_number,
                    column="cost_code_id",
                    issues=issues,
                ),
                label=_cell_text(raw_label),
                quantity=_cell_decimal(
                    raw_quantity,
                    sheet=_COST_LINE_SHEET,
                    row_number=row_number,
                    column="quantity",
                    issues=issues,
                ),
                unit_cost=_cell_decimal(
                    raw_unit_cost,
                    sheet=_COST_LINE_SHEET,
                    row_number=row_number,
                    column="unit_cost",
                    issues=issues,
                ),
                supply_status=_cell_text(raw_supply_status),
                planned_date=_cell_datetime(
                    raw_planned_date,
                    sheet=_COST_LINE_SHEET,
                    row_number=row_number,
                    column="planned_date",
                    issues=issues,
                ),
            )
        )
    return rows


def parse_estimate_reconciliation_workbook(content: bytes) -> ParsedReconciliationWorkbook:
    """Parse the three reconciliation sheets, or raise with every format problem at once.

    Read-only/data-only: values already computed by Excel are read verbatim,
    formulas are never evaluated by this parser.
    """
    issues: list[dict[str, object]] = []
    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except (BadZipFile, KeyError, OSError, ValueError) as exc:
        raise EstimateReconciliationFormatError(
            [
                {
                    "code": "INVALID_WORKBOOK",
                    "message": f"File is not a readable Excel workbook: {exc}",
                }
            ]
        ) from exc

    try:
        _check_headers(workbook, _TASK_SHEET, TASK_HEADERS, issues)
        _check_headers(workbook, _LABOR_SHEET, LABOR_HEADERS, issues)
        _check_headers(workbook, _COST_LINE_SHEET, COST_LINE_HEADERS, issues)
        if issues:
            raise EstimateReconciliationFormatError(issues)

        tasks = _parse_task_rows(workbook[_TASK_SHEET], issues)
        labor = _parse_labor_rows(workbook[_LABOR_SHEET], issues)
        cost_lines = _parse_cost_line_rows(workbook[_COST_LINE_SHEET], issues)
        if issues:
            raise EstimateReconciliationFormatError(issues)
    finally:
        workbook.close()

    return ParsedReconciliationWorkbook(tasks=tasks, labor=labor, cost_lines=cost_lines)


# --------------------------------------------------------------------------------------
# The revision reconciliation (E14-07c, #365)
#
# The same round trip, against the single tree of `wf_revision_node`. Everything
# above is the legacy parser, alive until E14-12 (#339) removes it with the three
# tables it reconciles -- and kept meanwhile because it is the only thing the
# workbook below can be compared against.
#
# Unlike its legacy half, this one is *not* a pure parser. The legacy diff/apply
# lives in `api/routes/estimates.py` because it needed route-layer helpers; the
# revision one needs none -- every rule it enforces is a domain rule, and the two
# things it does that the domain cannot are reading the referential (is this role an
# active labour role) and deciding which nodes the file no longer mentions. That is
# service work, and it sits next to the parser whose output it consumes rather than
# in a handler.
# --------------------------------------------------------------------------------------


@dataclass(frozen=True)
class RevisionTaskFileRow:
    """One row of the ``Tâches`` sheet of a revision workbook.

    ``node_id`` is the stable identifier, and the only one: where the legacy sheet
    carried a devis row id *and* the ``ms_task`` id it pointed at, a task is a node.
    Blank means "create this task".
    """

    row_number: int
    node_id: int | None
    #: The MS Project identity of the work item behind this node (E14-06), read back
    #: so that an edit to it can be *reported* rather than silently dropped. This
    #: import never writes it: see :func:`_stage_task_sheet`.
    external_uid: int | None
    parent_node_id: int | None
    position: int | None
    name: str | None
    is_milestone: bool


@dataclass(frozen=True)
class RevisionLaborFileRow:
    """One row of the ``MO`` sheet of a revision workbook."""

    row_number: int
    node_id: int | None
    parent_node_id: int | None
    position: int | None
    role_id: int | None
    cost_code_id: int | None
    label: str | None
    quantity: Decimal | None
    hours: Decimal | None
    comment: str | None


@dataclass(frozen=True)
class RevisionCostFileRow:
    """One row of the ``Non-MO`` sheet of a revision workbook."""

    row_number: int
    node_id: int | None
    parent_node_id: int | None
    position: int | None
    cost_category_id: int | None
    cost_code_id: int | None
    label: str | None
    quantity: Decimal | None
    unit_cost: Decimal | None
    supply_status: str | None
    planned_date: date | None


@dataclass(frozen=True)
class ParsedRevisionWorkbook:
    tasks: list[RevisionTaskFileRow]
    labor: list[RevisionLaborFileRow]
    cost_lines: list[RevisionCostFileRow]


def _cell_date(
    value: object,
    *,
    sheet: str,
    row_number: int,
    column: str,
    issues: list[dict[str, object]],
) -> date | None:
    """:func:`_cell_datetime`, narrowed to the calendar date the facet stores.

    ``wf_revision_cost_facet.planned_date`` is a ``Date`` where the legacy
    ``wf_estimate_cost_line.planned_date`` was a ``DateTime``, so the
    "compare at date-only precision" rule the legacy reconciliation had to carry
    (``_planned_dates_equal``, ``api/routes/estimates.py``) has nothing left to
    reconcile: the column and the exported cell now hold the same thing.
    """
    parsed = _cell_datetime(value, sheet=sheet, row_number=row_number, column=column, issues=issues)
    return None if parsed is None else parsed.date()


def _parse_revision_task_rows(
    sheet: Worksheet, issues: list[dict[str, object]]
) -> list[RevisionTaskFileRow]:
    rows: list[RevisionTaskFileRow] = []
    name = _TASK_SHEET
    for row_number, raw_values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        values = _pad(tuple(raw_values), len(REVISION_TASK_HEADERS))
        if _is_blank_row(values):
            continue
        (
            raw_node_id,
            _raw_work_item_id,
            raw_external_uid,
            raw_parent_node_id,
            raw_position,
            _raw_level,
            raw_name,
            raw_is_milestone,
        ) = values[: len(REVISION_TASK_HEADERS)]
        rows.append(
            RevisionTaskFileRow(
                row_number=row_number,
                node_id=_cell_int(
                    raw_node_id,
                    sheet=name,
                    row_number=row_number,
                    column="node_id",
                    issues=issues,
                ),
                external_uid=_cell_int(
                    raw_external_uid,
                    sheet=name,
                    row_number=row_number,
                    column="external_uid",
                    issues=issues,
                ),
                parent_node_id=_cell_int(
                    raw_parent_node_id,
                    sheet=name,
                    row_number=row_number,
                    column="parent_node_id",
                    issues=issues,
                ),
                position=_cell_int(
                    raw_position,
                    sheet=name,
                    row_number=row_number,
                    column="position",
                    issues=issues,
                ),
                name=_cell_text(raw_name),
                is_milestone=_cell_bool(raw_is_milestone),
            )
        )
    return rows


def _parse_revision_labor_rows(
    sheet: Worksheet, issues: list[dict[str, object]]
) -> list[RevisionLaborFileRow]:
    rows: list[RevisionLaborFileRow] = []
    name = _LABOR_SHEET
    for row_number, raw_values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        values = _pad(tuple(raw_values), len(REVISION_LABOR_HEADERS))
        if _is_blank_row(values):
            continue
        (
            raw_node_id,
            _raw_work_item_id,
            raw_parent_node_id,
            raw_position,
            _raw_bearing_node_id,
            _raw_bearing_name,
            raw_role_id,
            _raw_role_name,
            _raw_category_id,
            _raw_category_name,
            raw_cost_code_id,
            raw_label,
            raw_quantity,
            raw_hours,
            raw_comment,
        ) = values[: len(REVISION_LABOR_HEADERS)]
        rows.append(
            RevisionLaborFileRow(
                row_number=row_number,
                node_id=_cell_int(
                    raw_node_id,
                    sheet=name,
                    row_number=row_number,
                    column="node_id",
                    issues=issues,
                ),
                parent_node_id=_cell_int(
                    raw_parent_node_id,
                    sheet=name,
                    row_number=row_number,
                    column="parent_node_id",
                    issues=issues,
                ),
                position=_cell_int(
                    raw_position,
                    sheet=name,
                    row_number=row_number,
                    column="position",
                    issues=issues,
                ),
                role_id=_cell_int(
                    raw_role_id,
                    sheet=name,
                    row_number=row_number,
                    column="role_id",
                    issues=issues,
                ),
                cost_code_id=_cell_int(
                    raw_cost_code_id,
                    sheet=name,
                    row_number=row_number,
                    column="cost_code_id",
                    issues=issues,
                ),
                label=_cell_text(raw_label),
                quantity=_cell_decimal(
                    raw_quantity,
                    sheet=name,
                    row_number=row_number,
                    column="quantity",
                    issues=issues,
                ),
                hours=_cell_decimal(
                    raw_hours, sheet=name, row_number=row_number, column="hours", issues=issues
                ),
                comment=_cell_text(raw_comment),
            )
        )
    return rows


def _parse_revision_cost_rows(
    sheet: Worksheet, issues: list[dict[str, object]]
) -> list[RevisionCostFileRow]:
    rows: list[RevisionCostFileRow] = []
    name = _COST_LINE_SHEET
    for row_number, raw_values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        values = _pad(tuple(raw_values), len(REVISION_COST_LINE_HEADERS))
        if _is_blank_row(values):
            continue
        (
            raw_node_id,
            _raw_work_item_id,
            raw_parent_node_id,
            raw_position,
            _raw_bearing_node_id,
            _raw_bearing_name,
            _raw_cost_type_id,
            _raw_cost_type_code,
            raw_category_id,
            _raw_accounting_code,
            _raw_category_code,
            raw_cost_code_id,
            raw_label,
            raw_quantity,
            raw_unit_cost,
            _raw_purchase_cost,
            raw_supply_status,
            raw_planned_date,
        ) = values[: len(REVISION_COST_LINE_HEADERS)]
        rows.append(
            RevisionCostFileRow(
                row_number=row_number,
                node_id=_cell_int(
                    raw_node_id,
                    sheet=name,
                    row_number=row_number,
                    column="node_id",
                    issues=issues,
                ),
                parent_node_id=_cell_int(
                    raw_parent_node_id,
                    sheet=name,
                    row_number=row_number,
                    column="parent_node_id",
                    issues=issues,
                ),
                position=_cell_int(
                    raw_position,
                    sheet=name,
                    row_number=row_number,
                    column="position",
                    issues=issues,
                ),
                cost_category_id=_cell_int(
                    raw_category_id,
                    sheet=name,
                    row_number=row_number,
                    column="cost_category_id",
                    issues=issues,
                ),
                cost_code_id=_cell_int(
                    raw_cost_code_id,
                    sheet=name,
                    row_number=row_number,
                    column="cost_code_id",
                    issues=issues,
                ),
                label=_cell_text(raw_label),
                quantity=_cell_decimal(
                    raw_quantity,
                    sheet=name,
                    row_number=row_number,
                    column="quantity",
                    issues=issues,
                ),
                unit_cost=_cell_decimal(
                    raw_unit_cost,
                    sheet=name,
                    row_number=row_number,
                    column="unit_cost",
                    issues=issues,
                ),
                supply_status=_cell_text(raw_supply_status),
                planned_date=_cell_date(
                    raw_planned_date,
                    sheet=name,
                    row_number=row_number,
                    column="planned_date",
                    issues=issues,
                ),
            )
        )
    return rows


def parse_revision_reconciliation_workbook(content: bytes) -> ParsedRevisionWorkbook:
    """Parse the three sheets of a **revision** workbook, or raise with every problem.

    Word for word :func:`parse_estimate_reconciliation_workbook`'s contract -- the
    same three sheet names, the same "collect every format problem, never a partial
    parse", the same :class:`EstimateReconciliationFormatError` -- against the
    columns :mod:`waterfall.services.estimate_reconciliation_export` writes for a
    revision (:data:`~waterfall.services.estimate_reconciliation_export.REVISION_TASK_HEADERS`
    and its two siblings).

    Pure, like its legacy twin: bytes in, typed rows out, no session and no domain.
    What the file *means* against a given revision is :func:`reconcile_revision`.
    """
    issues: list[dict[str, object]] = []
    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except (BadZipFile, KeyError, OSError, ValueError) as exc:
        raise EstimateReconciliationFormatError(
            [
                {
                    "code": "INVALID_WORKBOOK",
                    "message": f"File is not a readable Excel workbook: {exc}",
                }
            ]
        ) from exc

    try:
        _check_headers(workbook, _TASK_SHEET, REVISION_TASK_HEADERS, issues)
        _check_headers(workbook, _LABOR_SHEET, REVISION_LABOR_HEADERS, issues)
        _check_headers(workbook, _COST_LINE_SHEET, REVISION_COST_LINE_HEADERS, issues)
        if issues:
            raise EstimateReconciliationFormatError(issues)

        tasks = _parse_revision_task_rows(workbook[_TASK_SHEET], issues)
        labor = _parse_revision_labor_rows(workbook[_LABOR_SHEET], issues)
        cost_lines = _parse_revision_cost_rows(workbook[_COST_LINE_SHEET], issues)
        if issues:
            raise EstimateReconciliationFormatError(issues)
    finally:
        workbook.close()

    return ParsedRevisionWorkbook(tasks=tasks, labor=labor, cost_lines=cost_lines)


# --------------------------------------------------------------------------------------
# What the file means against a revision: the diff, then the apply
# --------------------------------------------------------------------------------------


@dataclass(frozen=True)
class RevisionReconciliationIssue:
    """One structured entry of a :class:`RevisionReconciliationPlan` list.

    ``sheet``/``row`` point at the exact Excel cell a problem came from (``row`` is
    the 1-based Excel row number, header included); both are ``None`` for a problem
    that no single file row can be attributed to -- a deletion, by definition, has
    no row left in the file to point at.
    """

    code: str
    message: str
    sheet: str | None = None
    row: int | None = None


@dataclass(frozen=True)
class RevisionReconciliationPlan:
    """What a reconciliation file would do -- or did -- to a revision.

    ``blocking_issues`` non-empty means nothing was written and nothing will be:
    the counts below still describe the diff the file states, because they are what
    the user has to fix, and shrinking them would hide the very rows that need
    attention.

    ``cost_losses`` is Règle 3's safeguard, in the shape the rest of this API
    already publishes it (``POST .../nodes/delete``, the import diff): a row of the
    revision the file no longer mentions is a deletion, and a deletion that takes
    chiffrage away names it. Amounts are in euros at the cent, from the engine's
    published resolver, so the four callers quoting a loss -- the node deletion, the
    MS Project import diff, the MS Project import run and this round trip -- all
    quote the same figure (#368).
    """

    revision_id: int
    lock_version: int
    blocking_issues: tuple[RevisionReconciliationIssue, ...]
    warnings: tuple[RevisionReconciliationIssue, ...]
    tasks_to_create: int
    tasks_to_delete: tuple[int, ...]
    labor_to_create: int
    labor_to_update: tuple[int, ...]
    labor_to_delete: tuple[int, ...]
    non_labor_to_create: int
    non_labor_to_update: tuple[int, ...]
    non_labor_to_delete: tuple[int, ...]
    cost_losses: tuple[domain.CostLoss, ...]
    applied: bool


@dataclass(frozen=True)
class _Referential:
    """The referential rows the file's ids have to resolve against.

    Loaded once per reconciliation, keyed by the ids the file actually cites, so a
    workbook with two hundred MO rows still costs three queries and not six
    hundred.
    """

    #: Active labour roles only -- an inactive one, or one whose category is not of
    #: a labour cost type, is simply absent and therefore refused.
    labor_roles: dict[int, ResourceRole]
    #: Active non-labour categories, by id, with the ``kind`` of their cost type.
    non_labor_categories: dict[int, tuple[CostCategory, str]]
    #: Active cost codes of *this* project.
    cost_codes: set[int]


def _load_referential(
    db: Session,
    project_id: int,
    parsed: ParsedRevisionWorkbook,
    revision: domain.ProjectRevision,
) -> _Referential:
    """Resolve, in three queries, every referential id the file needs.

    ``revision`` is read for one thing only: a Non-MO row whose ``cost_category_id``
    cell is blank keeps the category the facet already carries, so that id has to be
    resolved too -- and reported as invalid if it has been deactivated since (the
    rule the legacy staging spelled out as E6-09 Finding Haute #3). Collecting it
    here is what lets :func:`_effective_category` answer without a query of its own.
    """
    role_ids = {row.role_id for row in parsed.labor if row.role_id is not None}
    labor_roles = (
        {
            role.id: role
            for role in db.query(ResourceRole)
            .join(CostCategory, ResourceRole.cost_category_id == CostCategory.id)
            .join(CostType, CostCategory.cost_type_id == CostType.id)
            .filter(ResourceRole.id.in_(role_ids))
            .filter(ResourceRole.is_active.is_(True))
            .filter(CostCategory.is_active.is_(True))
            .filter(CostType.is_active.is_(True))
            .filter(CostType.kind == CostTypeKind.LABOR)
            .all()
        }
        if role_ids
        else {}
    )
    category_ids = {row.cost_category_id for row in parsed.cost_lines if row.cost_category_id}
    category_ids.update(
        facet.cost_category_id
        for row in parsed.cost_lines
        if row.cost_category_id is None and row.node_id is not None
        for facet in (revision.cost_facets.get(row.node_id),)
        if facet is not None and facet.cost_category_id is not None
    )
    non_labor_categories = (
        {
            category.id: (category, kind)
            for category, kind in db.query(CostCategory, CostType.kind)
            .join(CostType, CostCategory.cost_type_id == CostType.id)
            .filter(CostCategory.id.in_(category_ids))
            .filter(CostCategory.is_active.is_(True))
            .filter(CostType.is_active.is_(True))
            .filter(CostType.kind != CostTypeKind.LABOR)
            .all()
        }
        if category_ids
        else {}
    )
    cost_code_ids = {
        row.cost_code_id
        for rows in (parsed.labor, parsed.cost_lines)
        for row in rows
        if row.cost_code_id is not None
    }
    cost_codes: set[int] = (
        {
            code_id
            for (code_id,) in db.query(ProjectCostCode.id)
            .filter(ProjectCostCode.id.in_(cost_code_ids))
            .filter(ProjectCostCode.project_id == project_id)
            .filter(ProjectCostCode.is_active.is_(True))
            .all()
        }
        if cost_code_ids
        else set()
    )
    return _Referential(
        labor_roles=labor_roles,
        non_labor_categories=non_labor_categories,
        cost_codes=cost_codes,
    )


class _Staging:
    """The accumulator every sheet stages into: issues, warnings and the plan.

    A small mutable object rather than six lists threaded through six signatures,
    which is what the legacy staging did and what made its helpers take seven
    parameters each.
    """

    def __init__(self, project: domain.Project, revision: domain.ProjectRevision) -> None:
        #: Read for one thing: ``external_uid`` is carried by the **work item**, not
        #: by the node, so the ``Tâches`` sheet cannot be compared without it.
        self.project = project
        self.revision = revision
        self.issues: list[RevisionReconciliationIssue] = []
        self.warnings: list[RevisionReconciliationIssue] = []
        self.seen_node_ids: set[int] = set()
        self.task_creates: list[RevisionTaskFileRow] = []
        self.labor_creates: list[RevisionLaborFileRow] = []
        self.labor_updates: list[RevisionLaborFileRow] = []
        self.cost_creates: list[RevisionCostFileRow] = []
        self.cost_updates: list[RevisionCostFileRow] = []

    def refuse(self, code: str, message: str, sheet: str, row: int) -> None:
        self.issues.append(
            RevisionReconciliationIssue(code=code, message=message, sheet=sheet, row=row)
        )

    def warn(self, code: str, message: str, sheet: str, row: int) -> None:
        self.warnings.append(
            RevisionReconciliationIssue(code=code, message=message, sheet=sheet, row=row)
        )

    def claim(self, node_id: int, kind: str, sheet: str, row: int) -> bool:
        """Check that ``node_id`` is a node of this revision, of the expected facet.

        ``kind`` is ``"task"`` or ``"cost"``. Returns ``False`` -- having already
        refused -- when the id is unknown, already claimed by another row, or
        carries the other facet. Claiming is what also feeds the deletion set: a
        node no row claims is a node the file dropped.
        """
        if node_id in self.seen_node_ids:
            self.refuse(
                "NODE_ID_DUPLICATED",
                f"node_id {node_id} appears on more than one row",
                sheet,
                row,
            )
            return False
        self.seen_node_ids.add(node_id)
        facets = self.revision.plan_facets if kind == "task" else self.revision.cost_facets
        if node_id not in self.revision.nodes:
            self.refuse(
                "NODE_ID_UNKNOWN",
                f"node_id {node_id} does not belong to this revision",
                sheet,
                row,
            )
            return False
        if node_id not in facets:
            self.refuse(
                "NODE_KIND_MISMATCH",
                f"node_id {node_id} does not carry a {kind} facet",
                sheet,
                row,
            )
            return False
        return True

    def parent(self, parent_node_id: int | None, sheet: str, row: int) -> bool:
        """Check a creation row's ``parent_node_id``: an existing task, or the root.

        Three refusals, all three of them refusals the **domain** would raise on the
        apply -- an unknown node, a task under a cost line (INV-11's facet
        placement), a child under a jalon (INV-27). Restated here rather than left
        to fire later because a preview that reported no blocking issue and a
        confirm that answered 400 would not be the same analysis, which is the one
        property the two endpoints promise.

        A node created by this same file cannot be a parent: it has no id yet, and
        the legacy import refused it for the same reason
        (``TASK_PARENT_UNRESOLVED``).
        """
        if parent_node_id is None:
            return True
        if parent_node_id not in self.revision.nodes:
            self.refuse(
                "NODE_PARENT_UNKNOWN",
                f"parent_node_id {parent_node_id} is not a node of this revision (a node "
                "created by this same file cannot be used as a parent)",
                sheet,
                row,
            )
            return False
        plan = self.revision.plan_facets.get(parent_node_id)
        if plan is None:
            self.refuse(
                "NODE_PARENT_IS_COST_LINE",
                f"parent_node_id {parent_node_id} is a cost line and holds no children",
                sheet,
                row,
            )
            return False
        if plan.is_milestone:
            self.refuse(
                "NODE_PARENT_IS_MILESTONE",
                f"parent_node_id {parent_node_id} is a milestone and holds no children",
                sheet,
                row,
            )
            return False
        return True

    def cost_code(
        self, cost_code_id: int | None, referential: _Referential, sheet: str, row: int
    ) -> bool:
        if cost_code_id is None or cost_code_id in referential.cost_codes:
            return True
        self.refuse(
            "COST_CODE_INVALID",
            f"cost_code_id {cost_code_id} does not belong to the project or is inactive",
            sheet,
            row,
        )
        return False


def _placement_changed(
    node: domain.RevisionNode, parent_node_id: int | None, position: int | None
) -> bool:
    return (parent_node_id is not None and parent_node_id != node.parent_id) or (
        position is not None and position != node.position
    )


def _stage_task_sheet(staging: _Staging, parsed: ParsedRevisionWorkbook) -> None:
    """The ``Tâches`` sheet: create or delete, never update.

    Exactly the capability the legacy reconciliation had, and deliberately not
    more: a difference on an existing row is a warning, not a write, because
    renaming and moving a task have a dedicated channel
    (``PATCH .../nodes/{id}/planning`` and ``POST .../nodes/move``) whose refusals
    -- a cycle, a jalon gaining a child, a position outside the sibling band -- are
    stated there once. An import quietly re-deriving the tree from a spreadsheet
    would be a second, weaker formulation of all three.
    """
    sheet = _TASK_SHEET
    for row in parsed.tasks:
        if row.node_id is not None:
            if not staging.claim(row.node_id, "task", sheet, row.row_number):
                continue
            node = staging.revision.nodes[row.node_id]
            facet = staging.revision.plan_facets[row.node_id]
            work_item = staging.project.work_items.get(node.work_item_id)
            changed = [
                name
                for name, changed_now in (
                    ("name", row.name is not None and row.name != facet.name),
                    ("is_milestone", row.is_milestone != facet.is_milestone),
                    ("placement", _placement_changed(node, row.parent_node_id, row.position)),
                    # `external_uid` is the MS Project identity of the work item
                    # (E14-06), not a derived echo like `work_item_id`/`level`: it is
                    # writable elsewhere, has a refusal of its own (INV-25,
                    # `REVISION_EXTERNAL_UID_CONFLICT`), and a user who retypes it here
                    # has every reason to expect something to happen. This import still
                    # will not write it -- rebinding a node to another MS Project task
                    # is the re-import's job, not a spreadsheet's -- but it says so
                    # instead of dropping the edit in silence.
                    (
                        "external_uid",
                        work_item is not None
                        and row.external_uid is not None
                        and row.external_uid != work_item.external_uid,
                    ),
                )
                if changed_now
            ]
            if changed:
                staging.warn(
                    "TASK_FIELD_CHANGE_IGNORED",
                    f"Ignored change(s) to: {', '.join(changed)} (renaming, moving or "
                    "re-identifying a task is not supported by this import; use the "
                    "revision tree endpoints, or the MS Project re-import for external_uid)",
                    sheet,
                    row.row_number,
                )
            continue
        if not row.name:
            staging.refuse(
                "TASK_NAME_REQUIRED", "name is required to create a task", sheet, row.row_number
            )
            continue
        if not staging.parent(row.parent_node_id, sheet, row.row_number):
            continue
        staging.task_creates.append(row)


def _stage_labor_row(
    staging: _Staging, referential: _Referential, row: RevisionLaborFileRow
) -> None:
    sheet = _LABOR_SHEET
    if row.label is None or row.quantity is None or row.hours is None:
        staging.refuse(
            "LABOR_ROW_INCOMPLETE",
            "label, quantity and hours are required on a MO row",
            sheet,
            row.row_number,
        )
        return
    if not staging.cost_code(row.cost_code_id, referential, sheet, row.row_number):
        return
    if row.node_id is None:
        if row.role_id is None or row.role_id not in referential.labor_roles:
            staging.refuse(
                "LABOR_ROLE_INVALID",
                f"role_id {row.role_id} must be an active labor role",
                sheet,
                row.row_number,
            )
            return
        if not staging.parent(row.parent_node_id, sheet, row.row_number):
            return
        staging.labor_creates.append(row)
        return
    facet = staging.revision.cost_facets[row.node_id]
    if row.role_id is not None and row.role_id != facet.role_id:
        staging.refuse(
            "LABOR_IDENTITY_CHANGE_REJECTED",
            "role_id cannot be changed on an existing MO line; delete it and add another",
            sheet,
            row.row_number,
        )
        return
    node = staging.revision.nodes[row.node_id]
    if _placement_changed(node, row.parent_node_id, row.position):
        staging.warn(
            "NODE_MOVE_IGNORED",
            "Ignored change to parent_node_id/position (moving a line is not supported by "
            "this import; use POST .../nodes/move)",
            sheet,
            row.row_number,
        )
    if (
        row.label == facet.label
        and row.quantity == facet.quantity
        and row.hours == facet.hours
        and (row.cost_code_id is None or row.cost_code_id == facet.cost_code_id)
        and row.comment == facet.comment
    ):
        # Unchanged against the stored facet: not staged, so an export followed by a
        # reimport with no edit at all produces an empty plan.
        return
    staging.labor_updates.append(row)


def _stage_labor_sheet(
    staging: _Staging, referential: _Referential, parsed: ParsedRevisionWorkbook
) -> None:
    for row in parsed.labor:
        if row.node_id is not None and not staging.claim(
            row.node_id, "cost", _LABOR_SHEET, row.row_number
        ):
            continue
        if row.node_id is not None and (
            staging.revision.cost_facets[row.node_id].nature is not domain.CostNature.LABOR
        ):
            staging.refuse(
                "NODE_KIND_MISMATCH",
                f"node_id {row.node_id} is not a labor cost line",
                _LABOR_SHEET,
                row.row_number,
            )
            continue
        _stage_labor_row(staging, referential, row)


def _effective_category(
    staging: _Staging, referential: _Referential, row: RevisionCostFileRow
) -> tuple[CostCategory, str] | None:
    """The (category, cost type kind) a Non-MO row lands on, or ``None`` once refused.

    A blank ``cost_category_id`` on an existing row means "unchanged", so the kind
    checked below is the *stored* line's own -- the rule the legacy staging had to
    spell out too (E6-09 Finding Haute #3): a category deactivated since the line
    was created must be reported here, not raised at apply time. Both ids are in
    :class:`_Referential` already, which is why this asks no question of its own.
    """
    category_id = row.cost_category_id
    if category_id is None and row.node_id is not None:
        category_id = staging.revision.cost_facets[row.node_id].cost_category_id
    if category_id is None:
        staging.refuse(
            "COST_LINE_ROW_INCOMPLETE",
            "cost_category_id is required on a Non-MO row",
            _COST_LINE_SHEET,
            row.row_number,
        )
        return None
    resolved = referential.non_labor_categories.get(category_id)
    if resolved is None:
        staging.refuse(
            "COST_LINE_CATEGORY_INVALID",
            f"cost_category_id {category_id} must be an active non-labor category",
            _COST_LINE_SHEET,
            row.row_number,
        )
        return None
    return resolved


def _stage_cost_row(staging: _Staging, referential: _Referential, row: RevisionCostFileRow) -> None:
    sheet = _COST_LINE_SHEET
    if row.label is None or row.quantity is None or row.unit_cost is None:
        staging.refuse(
            "COST_LINE_ROW_INCOMPLETE",
            "label, quantity and unit_cost are required on a Non-MO row",
            sheet,
            row.row_number,
        )
        return
    if not staging.cost_code(row.cost_code_id, referential, sheet, row.row_number):
        return
    resolved = _effective_category(staging, referential, row)
    if resolved is None:
        return
    _category, kind = resolved
    if row.supply_status is not None and kind != CostTypeKind.SUPPLY:
        staging.refuse(
            "COST_LINE_SUPPLY_STATUS_INVALID",
            "supply_status is only valid on a supply cost line",
            sheet,
            row.row_number,
        )
        return
    if row.supply_status is not None and row.supply_status not in {
        status.value for status in domain.SupplyStatus
    }:
        staging.refuse(
            "COST_LINE_SUPPLY_STATUS_INVALID",
            f"supply_status {row.supply_status!r} is not one of "
            f"{sorted(status.value for status in domain.SupplyStatus)}",
            sheet,
            row.row_number,
        )
        return
    if row.node_id is None:
        if not staging.parent(row.parent_node_id, sheet, row.row_number):
            return
        staging.cost_creates.append(row)
        return
    facet = staging.revision.cost_facets[row.node_id]
    node = staging.revision.nodes[row.node_id]
    if _placement_changed(node, row.parent_node_id, row.position):
        staging.warn(
            "NODE_MOVE_IGNORED",
            "Ignored change to parent_node_id/position (moving a line is not supported by "
            "this import; use POST .../nodes/move)",
            sheet,
            row.row_number,
        )
    if (
        row.label == facet.label
        and row.quantity == facet.quantity
        and row.unit_cost == facet.unit_cost
        and (row.cost_category_id is None or row.cost_category_id == facet.cost_category_id)
        and (row.cost_code_id is None or row.cost_code_id == facet.cost_code_id)
        and row.supply_status
        == (None if facet.supply_status is None else facet.supply_status.value)
        and row.planned_date == facet.planned_date
    ):
        return
    staging.cost_updates.append(row)


def _stage_cost_sheet(
    staging: _Staging, referential: _Referential, parsed: ParsedRevisionWorkbook
) -> None:
    for row in parsed.cost_lines:
        if row.node_id is not None and not staging.claim(
            row.node_id, "cost", _COST_LINE_SHEET, row.row_number
        ):
            continue
        if row.node_id is not None and (
            staging.revision.cost_facets[row.node_id].nature is domain.CostNature.LABOR
        ):
            staging.refuse(
                "NODE_KIND_MISMATCH",
                f"node_id {row.node_id} is a labor cost line",
                _COST_LINE_SHEET,
                row.row_number,
            )
            continue
        _stage_cost_row(staging, referential, row)


@dataclass(frozen=True)
class _Deletions:
    """The nodes the file no longer mentions, and the roots to hand the domain.

    ``roots`` is the subset whose parent is *not* itself doomed: deleting a node
    takes its whole subtree (INV-02), so handing the domain every doomed id would
    ask it twice for the same removal.
    """

    doomed: frozenset[int]
    roots: tuple[int, ...]


def _resolve_deletions(staging: _Staging) -> _Deletions:
    """Which nodes the file dropped, refusing the two shapes a drop must not take.

    A node of the revision that no row claims is a deletion -- that is the whole
    reconciliation contract, and it is what makes removing a line from the
    spreadsheet mean something. Two situations are refused rather than applied:

    * the file **keeps a child of a node it drops**. The cascade would take the
      child away anyway (INV-02), so applying it would delete a row the user
      explicitly left in the file -- the silent removal Règle 3 forbids;
    * the file **creates a node under a node it drops**, which would apply a write
      against a parent that no longer exists.
    """
    revision = staging.revision
    doomed = {node_id for node_id in revision.nodes if node_id not in staging.seen_node_ids}
    for node_id, node in sorted(revision.nodes.items()):
        if node_id in doomed or node.parent_id is None or node.parent_id not in doomed:
            continue
        staging.issues.append(
            RevisionReconciliationIssue(
                code="NODE_DELETE_CASCADE",
                message=(
                    f"node {node_id} is kept by the file but its parent {node.parent_id} is "
                    "not: removing the parent would take it away too (INV-02). Remove the "
                    "child from the file as well, or keep the parent."
                ),
            )
        )
    for sheet, rows in (
        (_TASK_SHEET, staging.task_creates),
        (_LABOR_SHEET, staging.labor_creates),
        (_COST_LINE_SHEET, staging.cost_creates),
    ):
        for row in rows:
            if row.parent_node_id is not None and row.parent_node_id in doomed:
                staging.refuse(
                    "NODE_PARENT_DELETED",
                    f"parent_node_id {row.parent_node_id} is removed by this same file",
                    sheet,
                    row.row_number,
                )
    roots = tuple(
        sorted(
            node_id
            for node_id in doomed
            if revision.nodes[node_id].parent_id is None
            or revision.nodes[node_id].parent_id not in doomed
        )
    )
    return _Deletions(doomed=frozenset(doomed), roots=roots)


def _supply_status(value: str | None) -> domain.SupplyStatus | None:
    return None if value is None else domain.SupplyStatus(value)


def _apply_reconciliation(
    loaded: LoadedRevision,
    staging: _Staging,
    referential: _Referential,
    deletions: _Deletions,
    amount_of: AmountResolver | None,
) -> None:
    """Run the staged plan as domain calls, deletions first and creations last.

    Deletions come first so that the positions they free are already free when a
    creation appends a sibling; creations come last because a node they would hang
    under has to have survived the deletions, which
    :func:`_resolve_deletions` has already refused the file for otherwise.

    Every refusal left is the domain's, and every one of them leaves the revision
    untouched provided the caller rolls back -- which
    :func:`waterfall.api.revision_errors.revision_operation` does.
    """
    project = loaded.project
    revision = loaded.revision
    now = datetime.now(UTC)
    if deletions.roots:
        # Non-`None` exactly when there is a root to delete, by construction: see
        # where `reconcile_revision` builds it.
        assert amount_of is not None
        domain.delete_nodes(project, revision, list(deletions.roots), amount_of=amount_of)
    for labor_row in staging.labor_updates:
        assert labor_row.node_id is not None
        domain.update_cost_facet(
            project,
            revision,
            labor_row.node_id,
            label=labor_row.label if labor_row.label is not None else domain.UNSET,
            quantity=labor_row.quantity if labor_row.quantity is not None else domain.UNSET,
            hours=labor_row.hours if labor_row.hours is not None else domain.UNSET,
            # A blank cost code cell means "unchanged", never "detach": the legacy
            # reconciliation settled that (E6-09, Critical finding) and a round trip
            # that silently unhooked every imputation would be unusable.
            #
            # `comment` -- and, on the sheet below, `supply_status`/`planned_date` --
            # go the *other* way on purpose: a blank cell clears them, which is why
            # they are passed straight through instead of falling back to `UNSET`.
            # The asymmetry is the difference between a *reference* and a free
            # attribute. Emptying a cost code means "this line is no longer imputed",
            # which is a decision with consequences on every aggregate by imputation,
            # and nobody clears a hundred of them on purpose by deleting a column;
            # emptying a comment means "there is no comment any more", and there is
            # no other way to say it -- a reconciliation file that could add a
            # comment but never remove one would be a one-way door.
            cost_code_id=(
                labor_row.cost_code_id if labor_row.cost_code_id is not None else domain.UNSET
            ),
            comment=labor_row.comment,
        )
    for cost_row in staging.cost_updates:
        assert cost_row.node_id is not None
        category = (
            None
            if cost_row.cost_category_id is None
            else referential.non_labor_categories[cost_row.cost_category_id][0]
        )
        domain.update_cost_facet(
            project,
            revision,
            cost_row.node_id,
            label=cost_row.label if cost_row.label is not None else domain.UNSET,
            quantity=cost_row.quantity if cost_row.quantity is not None else domain.UNSET,
            unit_cost=cost_row.unit_cost if cost_row.unit_cost is not None else domain.UNSET,
            # The cost type is never read off the file: it is a property of the
            # category, which is the single source of truth the removed create route
            # already used. Moving the line to another category moves it to that
            # category's type in the same call, so the two can never disagree.
            cost_type_id=domain.UNSET if category is None else category.cost_type_id,
            cost_category_id=domain.UNSET if category is None else category.id,
            supply_status=_supply_status(cost_row.supply_status),
            planned_date=cost_row.planned_date,
            cost_code_id=(
                cost_row.cost_code_id if cost_row.cost_code_id is not None else domain.UNSET
            ),
        )
    for task_row in staging.task_creates:
        assert task_row.name is not None
        domain.add_task(
            project,
            revision,
            name=task_row.name,
            parent_id=task_row.parent_node_id,
            is_milestone=task_row.is_milestone,
            now=now,
        )
    for labor_row in staging.labor_creates:
        assert labor_row.label is not None and labor_row.quantity is not None
        domain.add_cost_line(
            project,
            revision,
            nature=domain.CostNature.LABOR,
            label=labor_row.label,
            parent_id=labor_row.parent_node_id,
            quantity=labor_row.quantity,
            role_id=labor_row.role_id,
            hours=labor_row.hours,
            cost_code_id=labor_row.cost_code_id,
            comment=labor_row.comment,
            now=now,
        )
    for cost_row in staging.cost_creates:
        assert cost_row.label is not None and cost_row.quantity is not None
        assert cost_row.cost_category_id is not None
        category = referential.non_labor_categories[cost_row.cost_category_id][0]
        domain.add_cost_line(
            project,
            revision,
            nature=domain.CostNature.NON_LABOR,
            label=cost_row.label,
            parent_id=cost_row.parent_node_id,
            quantity=cost_row.quantity,
            cost_type_id=category.cost_type_id,
            cost_category_id=category.id,
            unit_cost=cost_row.unit_cost,
            supply_status=_supply_status(cost_row.supply_status),
            planned_date=cost_row.planned_date,
            cost_code_id=cost_row.cost_code_id,
            now=now,
        )


def _doomed_of(
    staging: _Staging, deletions: _Deletions, nature: domain.CostNature | None
) -> tuple[int, ...]:
    """The doomed node ids carrying a plan facet (``nature`` ``None``) or that nature."""
    revision = staging.revision
    if nature is None:
        return tuple(
            sorted(node_id for node_id in deletions.doomed if node_id in revision.plan_facets)
        )
    return tuple(
        sorted(
            node_id
            for node_id in deletions.doomed
            if node_id in revision.cost_facets and revision.cost_facets[node_id].nature is nature
        )
    )


def reconcile_revision(
    db: Session,
    revision_id: int,
    parsed: ParsedRevisionWorkbook,
    *,
    apply: bool,
    expected_lock_version: int | None = None,
) -> RevisionReconciliationPlan:
    """Diff a parsed revision workbook against ``revision_id``, and optionally apply it.

    One function for the two endpoints, called with ``apply=False`` by the preview
    and twice by the confirm (once unlocked as a precheck, once under the project
    lock with ``apply=True``), so that what a preview reports is by construction
    what a confirm on the same file does -- the property #70 established for the
    legacy round trip and which this one keeps.

    Writes nothing when ``apply`` is ``False``, and writes nothing when the staging
    found a blocking issue: the plan comes back with ``applied=False`` and the
    caller answers a 409 with it. Nothing here commits -- the caller owns the
    transaction, exactly as in :mod:`waterfall.services.revision_import`.

    ``expected_lock_version`` is the optimistic lock every other write of this API
    takes, offered here as a keyword because the file itself carries no version: the
    confirm passes the ``lock_version`` its own unlocked precheck read, and anything
    that wrote the revision in between is refused with the shared
    :class:`~waterfall.services.revision_tree.RevisionLockConflictError` -- 409
    ``REVISION_LOCK_CONFLICT``, nothing applied. It is checked on the **loaded**
    revision, before a single row is staged, so the refusal costs nothing and takes
    no domain call. Left ``None`` (the preview) nothing is compared, there being no
    write to protect.

    INV-03 is refused **up front** on the apply path, and not left to the first
    domain call: a file that turned out to describe the revision exactly as it
    already stands makes no domain call at all, and would otherwise answer 200 on a
    validated revision. The error raised is the domain's own
    :class:`~waterfall.domain.revision.ImmutableRevisionError`, so it comes back as
    the one ``REVISION_IMMUTABLE`` :mod:`waterfall.api.revision_errors` publishes
    for every other write of either facet -- there is no second table here.
    """
    loaded = load_revision(db, revision_id)
    revision = loaded.revision
    if expected_lock_version is not None and revision.lock_version != expected_lock_version:
        raise RevisionLockConflictError(revision_id, expected_lock_version, revision.lock_version)
    if apply and revision.status is not domain.RevisionStatus.DRAFT:
        raise domain.ImmutableRevisionError(
            f"Revision {revision_id} is {revision.status.value} and refuses every write "
            "(INV-03); copy it into a new draft before reconciling a workbook into it"
        )

    referential = _load_referential(db, revision.project_id, parsed, revision)
    staging = _Staging(loaded.project, revision)
    _stage_task_sheet(staging, parsed)
    _stage_labor_sheet(staging, referential, parsed)
    _stage_cost_sheet(staging, referential, parsed)
    deletions = _resolve_deletions(staging)

    # Priced before the first mutation and off the revision about to be mutated, so
    # the chiffrage the plan announces as lost is what the nodes were worth while
    # they were still there -- the same statement `revision_tree.delete_nodes` and
    # `revision_import.apply_import` make, in the same published unit (#368).
    #
    # And priced **only when the file drops something**: pricing walks every cost
    # node of the revision and resolves a rate per year per bearing task, and a
    # confirm runs this function twice. A file that deletes nothing -- the ordinary
    # case, an export edited in two cells and sent back -- has no loss to describe
    # and no deletion to make, so it pays none of that. `None` rather than a
    # neutral resolver on purpose: a wrong amount must not be *available* to be
    # quoted by mistake.
    amount_of: AmountResolver | None = (
        price_loaded_revision(db, loaded).published_amount_of if deletions.roots else None
    )
    cost_losses = (
        tuple(
            domain.describe_cost_losses(
                loaded.project, revision, list(deletions.roots), amount_of=amount_of
            )
        )
        if amount_of is not None
        else ()
    )

    plan = RevisionReconciliationPlan(
        revision_id=revision_id,
        lock_version=revision.lock_version,
        blocking_issues=tuple(staging.issues),
        warnings=tuple(staging.warnings),
        tasks_to_create=len(staging.task_creates),
        tasks_to_delete=_doomed_of(staging, deletions, None),
        labor_to_create=len(staging.labor_creates),
        labor_to_update=tuple(
            sorted(row.node_id for row in staging.labor_updates if row.node_id is not None)
        ),
        labor_to_delete=_doomed_of(staging, deletions, domain.CostNature.LABOR),
        non_labor_to_create=len(staging.cost_creates),
        non_labor_to_update=tuple(
            sorted(row.node_id for row in staging.cost_updates if row.node_id is not None)
        ),
        non_labor_to_delete=_doomed_of(staging, deletions, domain.CostNature.NON_LABOR),
        cost_losses=cost_losses,
        applied=False,
    )
    if staging.issues or not apply:
        return plan

    _apply_reconciliation(loaded, staging, referential, deletions, amount_of)
    save_revision(db, loaded)
    return replace(plan, lock_version=revision.lock_version, applied=True)
