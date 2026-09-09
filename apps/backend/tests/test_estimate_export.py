"""Tests for the human-readable devis Excel export (services/estimate_export.py)."""

from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any, cast

from openpyxl import load_workbook

from waterfall.db.session import get_session_factory
from waterfall.models.ms_core import MsProject
from waterfall.models.resources import (
    CostCategory,
    CostType,
    Estimate,
    EstimateCostLine,
    EstimateLine,
    ProjectCostCode,
    ResourceNode,
    ResourceRole,
)
from waterfall.services.estimate_calculation import UNASSIGNED_COST_CODE_LABEL
from waterfall.services.estimate_export import build_estimate_workbook


def _grid_sheet_records(sheet: Any) -> list[dict[str, Any]]:
    """Read the `Devis` grid sheet's data rows: unlike a plain header-first sheet, its
    header row is row 8 (rows 1-7 are the project/estimate info block written by
    `_write_header`), and the first blank `Nature` cell marks the start of the
    subtotal block below the data rows."""
    rows = list(sheet.iter_rows(min_row=8, values_only=True))
    headers = rows[0]
    return [dict(zip(headers, row, strict=True)) for row in rows[1:] if row[0] is not None]


def _seed_estimate_with_cost_codes() -> int:
    """Seed a project with a root cost code, one child cost code, one labor line
    attached to the child code, one non-labor line attached to the root code, and a
    second non-labor line with no `cost_code_id` at all -- exercising both the
    per-row column (issue #71 / E6-10) and the `by_cost_code` aggregate together.
    """
    session_factory = get_session_factory()
    with session_factory() as session:
        project = MsProject(
            owner_id=None,
            external_uid=None,
            source_version=2016,
            save_version_out=16,
            name="Export cost code test",
            schedule_from_start=True,
            start_date=datetime(2026, 1, 1, 8, 0, tzinfo=UTC),
            finish_date=datetime(2026, 12, 31, 18, 0, tzinfo=UTC),
            calendar_uid=1,
            minutes_per_day=480,
            minutes_per_week=2400,
            days_per_month=20,
            currency_code="EUR",
        )
        session.add(project)
        session.flush()

        root_code = ProjectCostCode(project_id=project.id, parent_id=None, code="1", name="Root")
        session.add(root_code)
        session.flush()
        child_code = ProjectCostCode(
            project_id=project.id, parent_id=root_code.id, code="1.1", name="Child"
        )
        session.add(child_code)
        session.flush()

        node = ResourceNode(code="DIRECTION", name="Direction")
        session.add(node)
        session.flush()

        labor_type = CostType(code="MO", name="Main d'oeuvre", kind="labor")
        session.add(labor_type)
        session.flush()
        labor_category = CostCategory(
            cost_type_id=labor_type.id, accounting_code="MO-DEV", name="Développement"
        )
        session.add(labor_category)
        session.flush()
        role = ResourceRole(node_id=node.id, cost_category_id=labor_category.id, name="Développeur")
        session.add(role)
        session.flush()

        supply_type = CostType(code="FOURNITURE", name="Fourniture", kind="supply")
        session.add(supply_type)
        session.flush()
        supply_category = CostCategory(
            cost_type_id=supply_type.id, accounting_code="FO-CABLE", name="Câbles"
        )
        session.add(supply_category)
        session.flush()

        estimate = Estimate(
            project_id=project.id,
            version_number=1,
            kind="initial",
            status="validated",
            currency_code="EUR",
        )
        session.add(estimate)
        session.flush()

        session.add_all(
            [
                EstimateLine(
                    estimate_id=estimate.id,
                    task_id=None,
                    role_id=role.id,
                    cost_code_id=child_code.id,
                    task_name="Développement",
                    role_code=role.name,
                    role_name=role.name,
                    accounting_code="MO-DEV",
                    year=2026,
                    quantity=Decimal("1"),
                    hours=Decimal("10"),
                    hourly_rate=Decimal("100"),
                    inflation_coefficient=Decimal("1"),
                    budget_cost=Decimal("1000.00"),
                ),
                EstimateCostLine(
                    estimate_id=estimate.id,
                    task_id=None,
                    cost_code_id=root_code.id,
                    cost_type_id=supply_type.id,
                    cost_category_id=supply_category.id,
                    cost_type_code="FOURNITURE",
                    accounting_code="FO-CABLE",
                    category_code="ACHAT",
                    label="Câbles réseau",
                    quantity=Decimal("5"),
                    unit_cost=Decimal("25.00"),
                    purchase_cost=Decimal("125.00"),
                ),
                EstimateCostLine(
                    estimate_id=estimate.id,
                    task_id=None,
                    cost_code_id=None,
                    cost_type_id=supply_type.id,
                    cost_category_id=supply_category.id,
                    cost_type_code="FOURNITURE",
                    accounting_code="FO-DIVERS",
                    category_code="ACHAT",
                    label="Fournitures diverses",
                    quantity=Decimal("2"),
                    unit_cost=Decimal("10.00"),
                    purchase_cost=Decimal("20.00"),
                ),
                # A real `POST .../validate` (see `calculate_estimate_lines`'s non-labor
                # step) snapshots one `EstimateLine` per non-labor `EstimateCostLine` too,
                # carrying that same `cost_code_id` -- `calculate_estimate_aggregates`
                # only ever reads `EstimateLine`, so these two mirror the cost lines
                # above to keep the aggregates sheet consistent with the grid sheet.
                EstimateLine(
                    estimate_id=estimate.id,
                    task_id=None,
                    role_id=None,
                    cost_code_id=root_code.id,
                    task_name="Câbles réseau",
                    role_code="",
                    role_name="",
                    accounting_code="FO-CABLE",
                    year=2026,
                    quantity=Decimal("5"),
                    hours=Decimal("0"),
                    hourly_rate=Decimal("0"),
                    inflation_coefficient=Decimal("1"),
                    budget_cost=Decimal("125.00"),
                ),
                EstimateLine(
                    estimate_id=estimate.id,
                    task_id=None,
                    role_id=None,
                    cost_code_id=None,
                    task_name="Fournitures diverses",
                    role_code="",
                    role_name="",
                    accounting_code="FO-DIVERS",
                    year=2026,
                    quantity=Decimal("2"),
                    hours=Decimal("0"),
                    hourly_rate=Decimal("0"),
                    inflation_coefficient=Decimal("1"),
                    budget_cost=Decimal("20.00"),
                ),
            ]
        )
        session.commit()
        return estimate.id


def test_grid_sheet_shows_cost_code_per_line_with_fallback_for_unassigned() -> None:
    """Every exported line (MO and non-MO) shows the cost-imputation code attached to
    it, and a line with no `cost_code_id` shows the shared fallback label rather than a
    broken/empty cell."""
    estimate_id = _seed_estimate_with_cost_codes()
    session_factory = get_session_factory()
    with session_factory() as session:
        estimate = session.get(Estimate, estimate_id)
        assert estimate is not None
        project = session.get(MsProject, estimate.project_id)
        assert project is not None
        workbook_bytes = build_estimate_workbook(session, project, estimate)

    workbook = load_workbook(BytesIO(workbook_bytes))
    grid_records = _grid_sheet_records(workbook["Devis"])

    by_label = {record["Libellé"]: record for record in grid_records}
    assert by_label["Développement (Développeur)"]["Code d'imputation"] == "1.1"
    assert by_label["Câbles réseau"]["Code d'imputation"] == "1"
    assert by_label["Fournitures diverses"]["Code d'imputation"] == UNASSIGNED_COST_CODE_LABEL


def test_aggregates_sheet_by_cost_code_subtotal_matches_line_sum() -> None:
    """The Agrégats sheet's "Par code d'imputation" section sums to the same total as
    the lines carrying each code in the grid sheet (issue #71 / E6-10)."""
    estimate_id = _seed_estimate_with_cost_codes()
    session_factory = get_session_factory()
    with session_factory() as session:
        estimate = session.get(Estimate, estimate_id)
        assert estimate is not None
        project = session.get(MsProject, estimate.project_id)
        assert project is not None
        workbook_bytes = build_estimate_workbook(session, project, estimate)

    workbook = load_workbook(BytesIO(workbook_bytes))
    aggregates_sheet = workbook["Agrégats"]

    rows = list(aggregates_sheet.iter_rows(values_only=True))
    header_row_index = next(
        index for index, row in enumerate(rows) if row[0] == "Code d'imputation"
    )
    by_cost_code: dict[str, float] = {}
    for row in rows[header_row_index + 1 :]:
        code, amount = row[0], row[1]
        if code is None:
            break
        by_cost_code[cast(str, code)] = cast(float, amount)

    assert by_cost_code == {
        "1.1": 1000.00,
        "1": 125.00,
        UNASSIGNED_COST_CODE_LABEL: 20.00,
    }
    total_unburdened_cost = aggregates_sheet["B5"].value
    assert sum(by_cost_code.values()) == total_unburdened_cost
