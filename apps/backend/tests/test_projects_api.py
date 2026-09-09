from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any, cast
from uuid import uuid4

from fastapi.testclient import TestClient
from httpx import Response
from openpyxl import load_workbook

from waterfall.db.session import get_session_factory
from waterfall.main import app
from waterfall.models.ms_core import MsProject, MsTask
from waterfall.models.planning import WfPlanning, WfPlanningTaskSnapshot
from waterfall.models.resources import (
    Calendar,
    CalendarWeekday,
    CostCategory,
    CostType,
    Estimate,
    EstimateCostLine,
    EstimateTaskRow,
    ProjectCostCode,
    ResourceNode,
    ResourceRole,
)
from waterfall.models.wf_core import WfTaskEnrichment


def _auth_headers(client: TestClient, email: str | None = None) -> dict[str, str]:
    email = email or f"projects.tester.{uuid4().hex}@example.com"
    password = "SuperSecret123!"

    register_response: Response = client.post(
        "/auth/register",
        json={"email": email, "password": password},
    )
    assert register_response.status_code == 201

    token_response: Response = client.post(
        "/auth/token",
        data={"username": email, "password": password},
    )
    assert token_response.status_code == 200
    token = token_response.json()["access_token"]

    return {"Authorization": f"Bearer {token}"}


def _current_user_id(client: TestClient, headers: dict[str, str]) -> int:
    response: Response = client.get("/auth/me", headers=headers)
    assert response.status_code == 200
    payload = cast(dict[str, Any], response.json())
    return cast(int, payload["id"])


def _seed_projects_and_tasks(owner_id: int) -> tuple[int, int]:
    session_factory = get_session_factory()
    with session_factory() as session:
        project = MsProject(
            owner_id=owner_id,
            external_uid=None,
            source_version=2016,
            save_version_out=16,
            name="Project API Test",
            schedule_from_start=True,
            start_date=datetime(2026, 1, 5, 8, 0, tzinfo=UTC),
            finish_date=datetime(2026, 1, 20, 18, 0, tzinfo=UTC),
            calendar_uid=1,
            minutes_per_day=480,
            minutes_per_week=2400,
            days_per_month=20,
            currency_code="EUR",
        )
        session.add(project)
        session.flush()

        # This helper bypasses POST /projects (which auto-creates the project's root
        # cost code -- see create_project), so it must uphold that same #62/E6-01
        # invariant itself: role-assignment/cost-line creation (E6-02/#63) requires
        # every project to always have an active root cost code to default to.
        session.add(
            ProjectCostCode(
                project_id=project.id,
                parent_id=None,
                code=f"PRJ-{project.id}",
                name=project.name,
            )
        )
        session.flush()

        task1 = MsTask(
            project_id=project.id,
            uid=1001,
            name="Task One",
            task_type=0,
            outline_number="1",
            outline_level=1,
            wbs="1",
            start_at=datetime(2026, 1, 5, 8, 0, tzinfo=UTC),
            finish_at=datetime(2026, 1, 6, 18, 0, tzinfo=UTC),
            duration_minutes=None,
            duration_format=None,
            work_minutes=None,
            percent_complete=20,
            is_summary=False,
            is_milestone=False,
            calendar_uid=1,
        )
        task2 = MsTask(
            project_id=project.id,
            uid=1002,
            name="Task Two",
            task_type=0,
            outline_number="2",
            outline_level=1,
            wbs="2",
            start_at=datetime(2026, 1, 7, 8, 0, tzinfo=UTC),
            finish_at=datetime(2026, 1, 10, 18, 0, tzinfo=UTC),
            duration_minutes=None,
            duration_format=None,
            work_minutes=None,
            percent_complete=0,
            is_summary=False,
            is_milestone=False,
            calendar_uid=1,
        )
        session.add_all([task1, task2])
        session.commit()
        return project.id, 2


def _seed_roles() -> tuple[int, int]:
    session_factory = get_session_factory()
    with session_factory() as session:
        root = ResourceNode(code="DIRECTION", name="Direction")
        child = ResourceNode(code="SERVICE", name="Service", parent_id=None)
        session.add_all([root, child])
        session.flush()
        child.parent_id = root.id

        labor_type = CostType(code="MO", name="Main d'oeuvre", kind="labor")
        supply_type = CostType(code="FOURNITURE", name="Fourniture", kind="supply")
        session.add_all([labor_type, supply_type])
        session.flush()
        labor_category = CostCategory(
            cost_type_id=labor_type.id,
            accounting_code="MO-DEV",
            category_code="IDEX",
            name="Développement",
        )
        supply_category = CostCategory(
            cost_type_id=supply_type.id,
            accounting_code="FO-CABLE",
            category_code="ACHAT",
            name="Câbles",
        )
        session.add_all([labor_category, supply_category])
        session.flush()
        labor_role = ResourceRole(
            node_id=child.id,
            cost_category_id=labor_category.id,
            name="Développeur",
        )
        supply_role = ResourceRole(
            node_id=child.id,
            cost_category_id=supply_category.id,
            name="Câble",
        )
        session.add_all([labor_role, supply_role])
        session.commit()
        return labor_role.id, supply_role.id


def _cost_category_id_for_role(role_id: int) -> int:
    session_factory = get_session_factory()
    with session_factory() as session:
        role = session.query(ResourceRole).filter(ResourceRole.id == role_id).one()
        return role.cost_category_id


def _list_cost_codes(
    client: TestClient, headers: dict[str, str], project_id: int
) -> list[dict[str, Any]]:
    response: Response = client.get(f"/projects/{project_id}/cost-codes", headers=headers)
    assert response.status_code == 200
    return cast(list[dict[str, Any]], response.json()["items"])


def _root_cost_code_id(client: TestClient, headers: dict[str, str], project_id: int) -> int:
    root = next(
        item for item in _list_cost_codes(client, headers, project_id) if item["parent_id"] is None
    )
    return cast(int, root["id"])


def test_get_projects_and_project_tasks() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, expected_tasks = _seed_projects_and_tasks(_current_user_id(client, headers))

        projects_response: Response = client.get("/projects", headers=headers)
        assert projects_response.status_code == 200
        raw_projects_body = projects_response.json()
        assert isinstance(raw_projects_body, dict)
        projects_payload = cast(list[dict[str, Any]], raw_projects_body["items"])
        assert raw_projects_body["total"] == len(projects_payload)
        assert len(projects_payload) >= 1
        assert any(project["id"] == project_id for project in projects_payload)

        project_response: Response = client.get(f"/projects/{project_id}", headers=headers)
        assert project_response.status_code == 200
        raw_project_payload = project_response.json()
        assert isinstance(raw_project_payload, dict)
        project_payload = cast(dict[str, Any], raw_project_payload)
        assert project_payload["id"] == project_id
        assert project_payload["name"] == "Project API Test"

        tasks_response: Response = client.get(
            f"/projects/{project_id}/tasks",
            headers=headers,
        )
        assert tasks_response.status_code == 200
        raw_tasks_body = tasks_response.json()
        assert isinstance(raw_tasks_body, dict)
        tasks_payload = cast(list[dict[str, Any]], raw_tasks_body["items"])
        assert len(tasks_payload) == expected_tasks
        assert tasks_payload[0]["project_id"] == project_id
        assert all("description" in task for task in tasks_payload)
        # E9-01 (#146): id_display is gone, replaced by a read-only row_number.
        # E9-02 (#147): row_number is now computed from the depth-first display order
        # (both seeded tasks are roots with no explicit position, so ties break on id,
        # i.e. creation order).
        assert "id_display" not in tasks_payload[0]
        row_number_by_uid = {task["uid"]: task["row_number"] for task in tasks_payload}
        assert row_number_by_uid == {1001: 1, 1002: 2}


def test_patch_task_description_and_read_back() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))

        patch_response: Response = client.patch(
            f"/projects/{project_id}/tasks/1001",
            json={"description": "Description enrichie depuis Waterfall"},
            headers=headers,
        )
        assert patch_response.status_code == 200
        patch_payload = patch_response.json()
        assert patch_payload["uid"] == 1001
        assert patch_payload["description"] == "Description enrichie depuis Waterfall"

        tasks_response: Response = client.get(
            f"/projects/{project_id}/tasks",
            headers=headers,
        )
        assert tasks_response.status_code == 200
        tasks_payload = cast(list[dict[str, Any]], tasks_response.json()["items"])

        task_by_uid = {task["uid"]: task for task in tasks_payload}
        assert task_by_uid[1001]["description"] == "Description enrichie depuis Waterfall"
        assert task_by_uid[1002]["description"] is None


def _seed_ms_task_group_with_eleven_children(owner_id: int) -> int:
    """Root summary task (uid 100) with 11 children (uid 101..111), each carrying a
    numeric ``position``/``outline_number`` pair that a plain lexicographic string sort
    on ``outline_number`` would mis-order (e.g. "1.10"/"1.11" sort before "1.2") --
    E9-02 (#147) regression: row_number must reflect the numeric depth-first order, not
    the ``outline_number`` string.
    """
    session_factory = get_session_factory()
    with session_factory() as session:
        project = MsProject(
            owner_id=owner_id,
            external_uid=None,
            source_version=2016,
            save_version_out=16,
            name="Eleven siblings",
            schedule_from_start=True,
            start_date=datetime(2026, 1, 5, 8, 0, tzinfo=UTC),
            finish_date=datetime(2026, 1, 20, 18, 0, tzinfo=UTC),
            calendar_uid=1,
            minutes_per_day=480,
            minutes_per_week=2400,
            days_per_month=20,
            currency_code="EUR",
        )
        session.add(project)
        session.flush()

        session.add(
            MsTask(
                project_id=project.id,
                uid=100,
                name="Group",
                task_type=0,
                outline_number="1",
                outline_level=1,
                is_summary=True,
                is_milestone=False,
            )
        )
        session.add_all(
            MsTask(
                project_id=project.id,
                uid=100 + child_position,
                name=f"Child {child_position}",
                task_type=0,
                parent_uid=100,
                position=child_position,
                outline_number=f"1.{child_position}",
                outline_level=2,
                is_summary=False,
                is_milestone=False,
            )
            for child_position in range(1, 12)
        )
        session.commit()
        return project.id


def test_project_tasks_row_number_ignores_lexicographic_outline_number_order() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id = _seed_ms_task_group_with_eleven_children(_current_user_id(client, headers))

        tasks_response: Response = client.get(
            f"/projects/{project_id}/tasks",
            headers=headers,
        )
        assert tasks_response.status_code == 200
        tasks_payload = cast(list[dict[str, Any]], tasks_response.json()["items"])

        assert [task["uid"] for task in tasks_payload] == list(range(100, 112))
        assert [task["row_number"] for task in tasks_payload] == list(range(1, 13))


def test_patch_task_description_not_found() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))

        response: Response = client.patch(
            f"/projects/{project_id}/tasks/999999",
            json={"description": "X"},
            headers=headers,
        )
        assert response.status_code == 404


def test_snapshot_only_task_rejects_legacy_assignment() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client, "projects.snapshot-assignment@example.com")
        project_id = cast(
            int,
            client.post(
                "/projects", json={"name": "Snapshot-only assignment"}, headers=headers
            ).json()["id"],
        )
        with get_session_factory()() as session:
            planning = WfPlanning(project_id=project_id, version_number=1, status="draft")
            session.add(planning)
            session.flush()
            session.add(
                WfPlanningTaskSnapshot(
                    planning_id=planning.id,
                    uid=9101,
                    name="Snapshot-only",
                    position=1,
                    is_summary=False,
                    is_milestone=False,
                )
            )
            session.query(MsProject).filter(MsProject.id == project_id).update(
                {MsProject.displayed_planning_id: planning.id}
            )
            session.commit()
        role_id, _ = _seed_roles()

        response = client.post(
            f"/projects/{project_id}/tasks/9101/role-assignments",
            json={"role_id": role_id, "quantity": 1, "hours": 1},
            headers=headers,
        )

        assert response.status_code == 409
        assert response.json()["detail"] == {"code": "GENERIC_ERROR"}


def test_delete_planning_task_referenced_by_cost_line_conflicts_without_mutation() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))

        session_factory = get_session_factory()
        with session_factory() as session:
            cost_type = CostType(code=f"MAT-{uuid4().hex[:8]}", name="Materiel")
            session.add(cost_type)
            session.flush()
            cost_category = CostCategory(
                cost_type_id=cost_type.id,
                accounting_code=f"MATCAT-{uuid4().hex[:8]}",
                name="Materiel",
            )
            session.add(cost_category)
            session.commit()
            cost_category_id = cost_category.id

        tasks_response: Response = client.get(f"/projects/{project_id}/tasks", headers=headers)
        task_id = next(
            task["id"]
            for task in cast(list[dict[str, Any]], tasks_response.json()["items"])
            if task["uid"] == 1001
        )

        # Creating a draft planning without a source copies the legacy MsTask
        # rows into snapshots with the same uid (see create_planning), so the
        # cost line above -- keyed off the legacy task_id -- also references
        # the new planning's task uid=1001 through is_task_referenced's
        # legacy-task bridge.
        planning_response: Response = client.post(
            f"/projects/{project_id}/plannings", json={}, headers=headers
        )
        assert planning_response.status_code == 201
        planning_id = cast(int, planning_response.json()["id"])

        estimate_response: Response = client.post(
            f"/projects/{project_id}/estimates",
            json={"kind": "initial", "currency_code": "EUR"},
            headers=headers,
        )
        estimate_id = cast(int, estimate_response.json()["id"])

        cost_line_response: Response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines",
            json={
                "task_id": task_id,
                "cost_category_id": cost_category_id,
                "label": "Materiel dedie",
                "quantity": 1,
                "unit_cost": 100,
            },
            headers=headers,
        )
        assert cost_line_response.status_code == 201

        delete_response: Response = client.post(
            f"/projects/{project_id}/plannings/{planning_id}/tasks/delete",
            json={"task_uids": [1001], "expected_revision": 0},
            headers=headers,
        )
        assert delete_response.status_code == 409
        detail = cast(dict[str, Any], delete_response.json())["detail"]
        assert detail["code"] == "TASK_REFERENCED"
        assert detail["task_uids"] == [1001]

        with session_factory() as session:
            remaining = (
                session.query(WfPlanningTaskSnapshot)
                .filter(WfPlanningTaskSnapshot.planning_id == planning_id)
                .filter(WfPlanningTaskSnapshot.uid == 1001)
                .first()
            )
        assert remaining is not None


def test_delete_planning_task_referenced_by_parent_task_row_conflicts_without_mutation() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))

        tasks_response: Response = client.get(f"/projects/{project_id}/tasks", headers=headers)
        task_rows = cast(list[dict[str, Any]], tasks_response.json()["items"])
        parent_task_id = cast(int, next(task["id"] for task in task_rows if task["uid"] == 1001))
        child_task_id = cast(int, next(task["id"] for task in task_rows if task["uid"] == 1002))

        planning_response: Response = client.post(
            f"/projects/{project_id}/plannings", json={}, headers=headers
        )
        assert planning_response.status_code == 201
        planning_id = cast(int, planning_response.json()["id"])

        estimate_response: Response = client.post(
            f"/projects/{project_id}/estimates",
            json={"kind": "initial", "currency_code": "EUR"},
            headers=headers,
        )
        assert estimate_response.status_code == 201
        estimate_id = cast(int, estimate_response.json()["id"])

        session_factory = get_session_factory()
        with session_factory() as session:
            row = (
                session.query(EstimateTaskRow)
                .filter(EstimateTaskRow.estimate_id == estimate_id)
                .filter(EstimateTaskRow.task_id == child_task_id)
                .first()
            )
            assert row is not None
            row.parent_task_id = parent_task_id
            session.commit()

        delete_response: Response = client.post(
            f"/projects/{project_id}/plannings/{planning_id}/tasks/delete",
            json={"task_uids": [1001], "expected_revision": 0},
            headers=headers,
        )
        assert delete_response.status_code == 409
        detail = cast(dict[str, Any], delete_response.json())["detail"]
        assert detail["code"] == "TASK_REFERENCED"
        assert detail["task_uids"] == [1001]


def test_estimate_aggregates_on_draft_are_zero() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))

        estimate_response: Response = client.post(
            f"/projects/{project_id}/estimates",
            json={"kind": "initial", "currency_code": "EUR"},
            headers=headers,
        )
        estimate_id = cast(int, estimate_response.json()["id"])

        response: Response = client.get(
            f"/projects/{project_id}/estimates/{estimate_id}/aggregates", headers=headers
        )
        assert response.status_code == 200
        payload = cast(dict[str, Any], response.json())
        assert float(payload["total_labor_cost"]) == 0
        assert float(payload["total_purchase_cost"]) == 0
        assert payload["by_category"] == {}


def test_estimate_aggregates_and_excel_export_after_validation() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))

        session_factory = get_session_factory()
        with session_factory() as session:
            cost_type = CostType(code=f"MAT-{uuid4().hex[:8]}", name="Materiel")
            session.add(cost_type)
            session.flush()
            cost_category = CostCategory(
                cost_type_id=cost_type.id,
                accounting_code=f"MATCAT-{uuid4().hex[:8]}",
                name="Materiel",
            )
            session.add(cost_category)
            session.commit()
            cost_category_id = cost_category.id

        estimate_response: Response = client.post(
            f"/projects/{project_id}/estimates",
            json={"kind": "initial", "currency_code": "EUR"},
            headers=headers,
        )
        estimate_id = cast(int, estimate_response.json()["id"])

        cost_line_response: Response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines",
            json={
                "cost_category_id": cost_category_id,
                "label": "Ordinateurs",
                "quantity": 2,
                "unit_cost": 900,
            },
            headers=headers,
        )
        assert cost_line_response.status_code == 201

        validate_response: Response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/validate", headers=headers
        )
        assert validate_response.status_code == 200

        aggregates_response: Response = client.get(
            f"/projects/{project_id}/estimates/{estimate_id}/aggregates", headers=headers
        )
        assert aggregates_response.status_code == 200
        aggregates_payload = cast(dict[str, Any], aggregates_response.json())
        assert float(aggregates_payload["total_purchase_cost"]) == 1800.0
        assert float(aggregates_payload["total_unburdened_cost"]) == 1800.0

        excel_response: Response = client.get(
            f"/projects/{project_id}/estimates/{estimate_id}/export.xlsx", headers=headers
        )
        assert excel_response.status_code == 200
        assert (
            excel_response.headers["content-type"]
            == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert excel_response.content[:2] == b"PK"


_TASK_SHEET_HEADERS = (
    "id",
    "task_id",
    "task_uid",
    "parent_task_id",
    "position",
    "task_name",
    "outline_number",
    "outline_level",
    "is_milestone",
)
_LABOR_SHEET_HEADERS = (
    "id",
    "task_id",
    "task_name",
    "role_id",
    "role_name",
    "cost_category_id",
    "cost_category_name",
    "cost_code_id",
    "quantity",
    "hours",
    "comment",
    "hors_perimetre_planning",
)
_NON_LABOR_SHEET_HEADERS = (
    "id",
    "task_id",
    "cost_type_id",
    "cost_type_code",
    "cost_category_id",
    "accounting_code",
    "category_code",
    "cost_code_id",
    "label",
    "quantity",
    "unit_cost",
    "purchase_cost",
    "supply_status",
    "planned_date",
)


def _sheet_records(sheet: Any) -> list[dict[str, Any]]:
    rows = list(sheet.iter_rows(values_only=True))
    headers = rows[0]
    return [dict(zip(headers, row, strict=True)) for row in rows[1:]]


def _fetch_task_id_by_uid(
    client: TestClient, headers: dict[str, str], project_id: int
) -> dict[int, int]:
    response: Response = client.get(f"/projects/{project_id}/tasks", headers=headers)
    assert response.status_code == 200
    return {
        task["uid"]: task["id"] for task in cast(list[dict[str, Any]], response.json()["items"])
    }


def _seed_reconciliation_fixture(client: TestClient, headers: dict[str, str]) -> dict[str, Any]:
    """Set up a project with one task-role assignment (MO) and one non-labor cost
    line, in a devis backed by a source planning -- E6-08 (#69) fixture shared by
    the reconciliation export tests below.
    """
    project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))

    planning_response: Response = client.post(
        f"/projects/{project_id}/plannings", json={}, headers=headers
    )
    assert planning_response.status_code == 201
    planning_id = cast(int, planning_response.json()["id"])

    # create_planning alone never touches displayed_planning_id (only
    # .../reference and .../display do); an estimate created below picks up
    # estimate.planning_id from project.displayed_planning_id, so it must be
    # set explicitly for the estimate to have a source planning at all.
    display_response: Response = client.post(
        f"/projects/{project_id}/plannings/{planning_id}/display", headers=headers
    )
    assert display_response.status_code == 200

    role_id, supply_role_id = _seed_roles()
    labor_category_id = _cost_category_id_for_role(role_id)
    supply_category_id = _cost_category_id_for_role(supply_role_id)

    task_id_by_uid = _fetch_task_id_by_uid(client, headers, project_id)

    assignment_response: Response = client.post(
        f"/projects/{project_id}/tasks/1001/role-assignments",
        json={"role_id": role_id, "quantity": 2, "hours": 10, "comment": "Dev senior"},
        headers=headers,
    )
    assert assignment_response.status_code == 201
    assignment_id = cast(int, assignment_response.json()["id"])

    estimate_response: Response = client.post(
        f"/projects/{project_id}/estimates",
        json={"kind": "initial", "currency_code": "EUR"},
        headers=headers,
    )
    assert estimate_response.status_code == 201
    estimate_id = cast(int, estimate_response.json()["id"])
    assert estimate_response.json()["planning_id"] == planning_id

    cost_line_response: Response = client.post(
        f"/projects/{project_id}/estimates/{estimate_id}/cost-lines",
        json={
            "task_id": task_id_by_uid[1002],
            "cost_category_id": supply_category_id,
            "label": "Cable reseau",
            "quantity": 5,
            "unit_cost": 12.5,
        },
        headers=headers,
    )
    assert cost_line_response.status_code == 201
    cost_line_id = cast(int, cost_line_response.json()["id"])

    task_rows_response: Response = client.get(
        f"/projects/{project_id}/estimates/{estimate_id}/task-rows", headers=headers
    )
    assert task_rows_response.status_code == 200
    task_row_by_task_id = {
        row["task_id"]: row["id"]
        for row in cast(list[dict[str, Any]], task_rows_response.json()["items"])
    }

    return {
        "project_id": project_id,
        "estimate_id": estimate_id,
        "task1_id": task_id_by_uid[1001],
        "task2_id": task_id_by_uid[1002],
        "task_row_id_for_task1": task_row_by_task_id[task_id_by_uid[1001]],
        "role_id": role_id,
        "labor_category_id": labor_category_id,
        "assignment_id": assignment_id,
        "supply_category_id": supply_category_id,
        "cost_line_id": cost_line_id,
    }


def test_estimate_reconciliation_export_contains_tasks_labor_and_non_labor_sheets() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        fixture = _seed_reconciliation_fixture(client, headers)

        response: Response = client.get(
            "/projects/"
            f"{fixture['project_id']}/estimates/{fixture['estimate_id']}"
            "/export-reconciliation.xlsx",
            headers=headers,
        )
        assert response.status_code == 200
        assert (
            response.headers["content-type"]
            == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert response.content[:2] == b"PK"

        workbook = load_workbook(BytesIO(cast(bytes, response.content)))
        assert workbook.sheetnames == ["Tâches", "MO", "Non-MO"]

        tasks_sheet = workbook["Tâches"]
        assert tuple(next(tasks_sheet.iter_rows(values_only=True))) == _TASK_SHEET_HEADERS
        task_records = _sheet_records(tasks_sheet)
        assert len(task_records) == 2
        task1_record = next(
            record for record in task_records if record["task_id"] == fixture["task1_id"]
        )
        assert task1_record["id"] == fixture["task_row_id_for_task1"]
        assert task1_record["task_uid"] == 1001
        assert task1_record["task_name"] == "Task One"
        assert task1_record["is_milestone"] is False

        labor_sheet = workbook["MO"]
        assert tuple(next(labor_sheet.iter_rows(values_only=True))) == _LABOR_SHEET_HEADERS
        labor_records = _sheet_records(labor_sheet)
        assert len(labor_records) == 1
        labor_record = labor_records[0]
        assert labor_record["id"] == fixture["assignment_id"]
        assert labor_record["task_id"] == fixture["task1_id"]
        assert labor_record["task_name"] == "Task One"
        assert labor_record["role_id"] == fixture["role_id"]
        assert labor_record["role_name"] == "Développeur"
        assert labor_record["cost_category_id"] == fixture["labor_category_id"]
        assert labor_record["quantity"] == 2
        assert labor_record["hours"] == 10
        assert labor_record["comment"] == "Dev senior"
        assert labor_record["hors_perimetre_planning"] is False

        non_labor_sheet = workbook["Non-MO"]
        assert tuple(next(non_labor_sheet.iter_rows(values_only=True))) == _NON_LABOR_SHEET_HEADERS
        non_labor_records = _sheet_records(non_labor_sheet)
        assert len(non_labor_records) == 1
        non_labor_record = non_labor_records[0]
        assert non_labor_record["id"] == fixture["cost_line_id"]
        assert non_labor_record["task_id"] == fixture["task2_id"]
        assert non_labor_record["cost_category_id"] == fixture["supply_category_id"]
        assert non_labor_record["label"] == "Cable reseau"
        assert non_labor_record["quantity"] == 5
        assert non_labor_record["unit_cost"] == 12.5
        assert non_labor_record["purchase_cost"] == 62.5
        assert non_labor_record["planned_date"] is None


def test_estimate_reconciliation_export_row_ids_are_stable_across_exports() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        fixture = _seed_reconciliation_fixture(client, headers)
        url = (
            "/projects/"
            f"{fixture['project_id']}/estimates/{fixture['estimate_id']}"
            "/export-reconciliation.xlsx"
        )

        first_response: Response = client.get(url, headers=headers)
        second_response: Response = client.get(url, headers=headers)
        assert first_response.status_code == 200
        assert second_response.status_code == 200

        first_workbook = load_workbook(BytesIO(cast(bytes, first_response.content)))
        second_workbook = load_workbook(BytesIO(cast(bytes, second_response.content)))
        for sheet_name in ("Tâches", "MO", "Non-MO"):
            first_ids = [record["id"] for record in _sheet_records(first_workbook[sheet_name])]
            second_ids = [record["id"] for record in _sheet_records(second_workbook[sheet_name])]
            assert first_ids == second_ids
            assert first_ids  # each sheet has at least the one seeded row


def test_estimate_reconciliation_export_without_source_planning_does_not_crash() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))
        role_id, _ = _seed_roles()

        assignment_response: Response = client.post(
            f"/projects/{project_id}/tasks/1001/role-assignments",
            json={"role_id": role_id, "quantity": 1, "hours": 1},
            headers=headers,
        )
        assert assignment_response.status_code == 201

        estimate_response: Response = client.post(
            f"/projects/{project_id}/estimates",
            json={"kind": "initial", "currency_code": "EUR"},
            headers=headers,
        )
        assert estimate_response.status_code == 201
        estimate_id = cast(int, estimate_response.json()["id"])
        assert estimate_response.json()["planning_id"] is None

        response: Response = client.get(
            f"/projects/{project_id}/estimates/{estimate_id}/export-reconciliation.xlsx",
            headers=headers,
        )
        assert response.status_code == 200

        workbook = load_workbook(BytesIO(cast(bytes, response.content)))
        assert len(_sheet_records(workbook["Tâches"])) == 2
        labor_records = _sheet_records(workbook["MO"])
        assert len(labor_records) == 1
        # No source planning at all: nothing can be "outside" a snapshot that doesn't exist.
        assert labor_records[0]["hors_perimetre_planning"] is False


def test_estimate_reconciliation_export_excludes_labor_cost_line_defensively() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))

        estimate_response: Response = client.post(
            f"/projects/{project_id}/estimates",
            json={"kind": "initial", "currency_code": "EUR"},
            headers=headers,
        )
        estimate_id = cast(int, estimate_response.json()["id"])

        session_factory = get_session_factory()
        with session_factory() as session:
            labor_type = CostType(code=f"MO-{uuid4().hex[:8]}", name="Main d'oeuvre", kind="labor")
            session.add(labor_type)
            session.flush()
            labor_category = CostCategory(
                cost_type_id=labor_type.id,
                accounting_code=f"MOCAT-{uuid4().hex[:8]}",
                name="Main d'oeuvre",
            )
            session.add(labor_category)
            session.flush()
            # Direct ORM insert: the API (get_non_labor_category_or_400) never lets a
            # labor cost type through create_estimate_cost_line, so this defensive
            # scenario can only be reproduced by bypassing the route layer.
            stray_line = EstimateCostLine(
                estimate_id=estimate_id,
                task_id=None,
                cost_code_id=None,
                cost_type_id=labor_type.id,
                cost_category_id=labor_category.id,
                cost_type_code=labor_type.code,
                accounting_code=labor_category.accounting_code,
                category_code=None,
                label="Ligne MO egaree",
                quantity=Decimal("1"),
                unit_cost=Decimal("100"),
                purchase_cost=Decimal("100"),
                supply_status=None,
                planned_date=None,
            )
            session.add(stray_line)
            session.commit()

        response: Response = client.get(
            f"/projects/{project_id}/estimates/{estimate_id}/export-reconciliation.xlsx",
            headers=headers,
        )
        assert response.status_code == 200

        workbook = load_workbook(BytesIO(cast(bytes, response.content)))
        non_labor_records = _sheet_records(workbook["Non-MO"])
        assert all(record["label"] != "Ligne MO egaree" for record in non_labor_records)


def test_estimate_reconciliation_export_flags_assignment_outside_planning_snapshot() -> None:
    """PR review finding Moyenne #1 (#69): a role assignment whose task was added
    to the project *after* the estimate's source planning was snapshotted (or
    whose task otherwise never made it into that snapshot) must still appear on
    the "MO" sheet -- never silently dropped -- flagged via
    ``hors_perimetre_planning=True`` so a reconciliation user can see it won't be
    reprised automatically on a future reimport (E6-09)."""
    with TestClient(app) as client:
        headers = _auth_headers(client)
        fixture = _seed_reconciliation_fixture(client, headers)
        project_id = cast(int, fixture["project_id"])

        # Added after the planning (and its snapshot) already exist, so this task's
        # uid is absent from `WfPlanningTaskSnapshot` for the estimate's source planning.
        session_factory = get_session_factory()
        with session_factory() as session:
            outside_task = MsTask(
                project_id=project_id,
                uid=1003,
                name="Task Outside Snapshot",
                task_type=0,
                outline_number="3",
                outline_level=1,
                wbs="3",
                start_at=datetime(2026, 1, 11, 8, 0, tzinfo=UTC),
                finish_at=datetime(2026, 1, 12, 18, 0, tzinfo=UTC),
                duration_minutes=None,
                duration_format=None,
                work_minutes=None,
                percent_complete=0,
                is_summary=False,
                is_milestone=False,
                calendar_uid=1,
            )
            session.add(outside_task)
            session.commit()

        outside_assignment_response: Response = client.post(
            f"/projects/{project_id}/tasks/1003/role-assignments",
            json={"role_id": fixture["role_id"], "quantity": 1, "hours": 5},
            headers=headers,
        )
        assert outside_assignment_response.status_code == 201
        outside_assignment_id = cast(int, outside_assignment_response.json()["id"])

        response: Response = client.get(
            f"/projects/{project_id}/estimates/{fixture['estimate_id']}/export-reconciliation.xlsx",
            headers=headers,
        )
        assert response.status_code == 200

        workbook = load_workbook(BytesIO(cast(bytes, response.content)))
        labor_records = _sheet_records(workbook["MO"])
        assert len(labor_records) == 2

        outside_record = next(
            record for record in labor_records if record["id"] == outside_assignment_id
        )
        assert outside_record["task_name"] == "Task Outside Snapshot"
        assert outside_record["hors_perimetre_planning"] is True

        in_scope_record = next(
            record for record in labor_records if record["id"] == fixture["assignment_id"]
        )
        assert in_scope_record["hors_perimetre_planning"] is False


def test_estimate_reconciliation_export_does_not_leak_role_assignments_across_projects() -> None:
    """Non-regression test for PR review finding Moyenne #2 (#69): TaskRoleAssignment
    is a project-wide table (not scoped to a single estimate), so a regression in
    `_scoped_task_role_assignments`'s `MsTask.project_id == project.id` filter
    could silently pull another project's labor assignments into this one's
    reconciliation export."""
    with TestClient(app) as client:
        headers = _auth_headers(client)
        owner_id = _current_user_id(client, headers)
        project_a_id, _ = _seed_projects_and_tasks(owner_id)
        project_b_id, _ = _seed_projects_and_tasks(owner_id)
        role_id, _ = _seed_roles()

        leaking_assignment_response: Response = client.post(
            f"/projects/{project_b_id}/tasks/1001/role-assignments",
            json={
                "role_id": role_id,
                "quantity": 1,
                "hours": 1,
                "comment": "PROJECT_B_ONLY_MUST_NOT_LEAK",
            },
            headers=headers,
        )
        assert leaking_assignment_response.status_code == 201

        estimate_response: Response = client.post(
            f"/projects/{project_a_id}/estimates",
            json={"kind": "initial", "currency_code": "EUR"},
            headers=headers,
        )
        assert estimate_response.status_code == 201
        estimate_id = cast(int, estimate_response.json()["id"])

        response: Response = client.get(
            f"/projects/{project_a_id}/estimates/{estimate_id}/export-reconciliation.xlsx",
            headers=headers,
        )
        assert response.status_code == 200

        workbook = load_workbook(BytesIO(cast(bytes, response.content)))
        labor_records = _sheet_records(workbook["MO"])
        assert all(record["comment"] != "PROJECT_B_ONLY_MUST_NOT_LEAK" for record in labor_records)
        assert labor_records == []


def test_estimate_reconciliation_export_project_not_found() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        response: Response = client.get(
            "/projects/999999/estimates/1/export-reconciliation.xlsx", headers=headers
        )
        assert response.status_code == 404


def test_estimate_reconciliation_export_estimate_not_found() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))

        response: Response = client.get(
            f"/projects/{project_id}/estimates/999999/export-reconciliation.xlsx",
            headers=headers,
        )
        assert response.status_code == 404


def test_get_project_not_found() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        response: Response = client.get("/projects/999999", headers=headers)
        assert response.status_code == 404


def test_get_projects_requires_auth() -> None:
    with TestClient(app) as client:
        response: Response = client.get("/projects")
        assert response.status_code == 401


def test_patch_project_name() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))

        response: Response = client.patch(
            f"/projects/{project_id}",
            json={"name": "Projet Renomme"},
            headers=headers,
        )
        assert response.status_code == 200
        payload = cast(dict[str, Any], response.json())
        assert payload["id"] == project_id
        assert payload["name"] == "Projet Renomme"


def test_create_project_with_code_and_description() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)

        response: Response = client.post(
            "/projects",
            json={"name": "Projet Code", "code": "PRJ-001", "short_description": "Résumé court"},
            headers=headers,
        )
        assert response.status_code == 201
        payload = cast(dict[str, Any], response.json())
        assert payload["code"] == "PRJ-001"
        assert payload["short_description"] == "Résumé court"


def test_create_project_exposes_calendar_defaults() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)

        create_response: Response = client.post(
            "/projects",
            json={"name": "Projet Calendrier"},
            headers=headers,
        )
        assert create_response.status_code == 201
        create_payload = cast(dict[str, Any], create_response.json())
        assert create_payload["minutes_per_day"] == 480
        assert create_payload["minutes_per_week"] == 2400
        assert create_payload["days_per_month"] == 20

        project_id = create_payload["id"]
        get_response: Response = client.get(f"/projects/{project_id}", headers=headers)
        assert get_response.status_code == 200
        get_payload = cast(dict[str, Any], get_response.json())
        assert get_payload["minutes_per_day"] == 480
        assert get_payload["minutes_per_week"] == 2400
        assert get_payload["days_per_month"] == 20


def test_patch_project_code_and_description_partially() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))

        first_response: Response = client.patch(
            f"/projects/{project_id}",
            json={"code": "PRJ-042"},
            headers=headers,
        )
        assert first_response.status_code == 200
        first_payload = cast(dict[str, Any], first_response.json())
        assert first_payload["code"] == "PRJ-042"
        assert first_payload["short_description"] is None
        assert first_payload["name"] == "Project API Test"

        second_response: Response = client.patch(
            f"/projects/{project_id}",
            json={"short_description": "Nouvelle description"},
            headers=headers,
        )
        assert second_response.status_code == 200
        second_payload = cast(dict[str, Any], second_response.json())
        assert second_payload["code"] == "PRJ-042"
        assert second_payload["short_description"] == "Nouvelle description"


def test_delete_project_cascades_related_data() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))

        session_factory = get_session_factory()
        with session_factory() as session:
            enrichment = WfTaskEnrichment(
                project_id=project_id,
                task_uid=1001,
                description="A supprimer",
                created_at=datetime.now(UTC),
                updated_at=datetime.now(UTC),
            )
            session.add(enrichment)
            session.commit()

            cost_type = CostType(code=f"MAT-{uuid4().hex[:8]}", name="Materiel")
            session.add(cost_type)
            session.flush()
            cost_category = CostCategory(
                cost_type_id=cost_type.id,
                accounting_code=f"MATCAT-{uuid4().hex[:8]}",
                name="Materiel",
            )
            session.add(cost_category)
            session.commit()
            cost_category_id = cost_category.id

        estimate_response: Response = client.post(
            f"/projects/{project_id}/estimates",
            json={"kind": "initial", "currency_code": "EUR"},
            headers=headers,
        )
        assert estimate_response.status_code == 201
        estimate_id = cast(int, estimate_response.json()["id"])

        cost_line_response: Response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines",
            json={
                "cost_category_id": cost_category_id,
                "label": "Ordinateurs",
                "quantity": 2,
                "unit_cost": 900,
            },
            headers=headers,
        )
        assert cost_line_response.status_code == 201

        response: Response = client.delete(f"/projects/{project_id}", headers=headers)
        assert response.status_code == 204

        project_response: Response = client.get(f"/projects/{project_id}", headers=headers)
        assert project_response.status_code == 404

        tasks_response: Response = client.get(f"/projects/{project_id}/tasks", headers=headers)
        assert tasks_response.status_code == 404


def test_delete_project_with_displayed_planning_snapshot() -> None:
    # Regression test: fk_ms_project_displayed_planning rejected the raw `DELETE FROM
    # wf_planning` issued by delete_project's bulk delete (synchronize_session=False) because the
    # session has autoflush disabled (db/session.py) and the preceding
    # `project.displayed_planning_id = None` ORM-level update was never flushed before that raw
    # delete ran. _seed_projects_and_tasks (used by test_delete_project_cascades_related_data
    # above) never creates a WfPlanning row, so it never exercised this path -- this test targets
    # exactly the wf_planning/wf_planning_task_snapshot model instead.
    with TestClient(app) as client:
        headers = _auth_headers(client, "projects.delete-displayed-planning@example.com")
        project_id = cast(
            int,
            client.post(
                "/projects", json={"name": "Delete with displayed planning"}, headers=headers
            ).json()["id"],
        )
        with get_session_factory()() as session:
            planning = WfPlanning(project_id=project_id, version_number=1, status="draft")
            session.add(planning)
            session.flush()
            session.add(
                WfPlanningTaskSnapshot(
                    planning_id=planning.id,
                    uid=1,
                    name="Snapshot task",
                    position=1,
                    is_summary=False,
                    is_milestone=False,
                )
            )
            session.query(MsProject).filter(MsProject.id == project_id).update(
                {MsProject.displayed_planning_id: planning.id}
            )
            session.commit()

        response: Response = client.delete(f"/projects/{project_id}", headers=headers)
        assert response.status_code == 204

        project_response: Response = client.get(f"/projects/{project_id}", headers=headers)
        assert project_response.status_code == 404


def test_creates_new_version_from_a_hierarchical_validated_planning() -> None:
    # Regression for #103: cloning a validated planning via source_planning_id used to
    # insert every cloned WfPlanningTaskSnapshot with its real parent_uid already set in
    # one batch. parent_uid is a composite self-reference onto (planning_id, uid) within
    # that same batch, so a child snapshot landing before its parent in the source
    # query's result violated the FK.
    #
    # The unfiltered `SELECT ... WHERE planning_id = X` in create_planning has no
    # ORDER BY, and SQLite favours the (planning_id, uid) index backing the composite
    # FK/uniqueness constraint for such a small table, returning rows in ascending
    # *uid* order rather than insertion order (see test_planning_clone_postgres.py for
    # why PostgreSQL needs a different reproduction: it plans that same query as a
    # sequential/heap scan for a table this size, returning insertion order instead).
    # Root is therefore given the highest uid and grandchild the lowest, decoupling
    # "hierarchy depth" from "uid value" so the query reliably returns
    # grandchild/child/root -- child-before-parent -- and reproduces the bug here,
    # while every row is still perfectly valid to insert in straightforward
    # root-then-child-then-grandchild order (each parent already exists by the time
    # its child references it).
    with TestClient(app) as client:
        headers = _auth_headers(client, "projects.hierarchical-clone@example.com")
        owner_id = _current_user_id(client, headers)

        session_factory = get_session_factory()
        with session_factory() as session:
            project = MsProject(
                owner_id=owner_id,
                source_version=2016,
                save_version_out=16,
                name="Hierarchical clone source",
                schedule_from_start=True,
                start_date=datetime(2026, 1, 5, tzinfo=UTC),
                finish_date=datetime(2026, 1, 20, tzinfo=UTC),
                minutes_per_day=480,
                minutes_per_week=2400,
                days_per_month=20,
            )
            session.add(project)
            session.flush()
            project_id = project.id

            planning = WfPlanning(project_id=project_id, version_number=1, status="draft")
            session.add(planning)
            session.flush()
            planning_id = planning.id

            session.add(
                WfPlanningTaskSnapshot(
                    planning_id=planning_id,
                    uid=100,
                    parent_uid=None,
                    name="Root",
                    position=1,
                    is_summary=True,
                    is_milestone=False,
                )
            )
            session.add(
                WfPlanningTaskSnapshot(
                    planning_id=planning_id,
                    uid=2,
                    parent_uid=100,
                    name="Child",
                    position=1,
                    is_summary=False,
                    is_milestone=False,
                )
            )
            session.add(
                WfPlanningTaskSnapshot(
                    planning_id=planning_id,
                    uid=1,
                    parent_uid=2,
                    name="Grandchild",
                    position=1,
                    is_summary=False,
                    is_milestone=False,
                )
            )
            project.displayed_planning_id = planning_id
            session.commit()

        validate_response = client.post(
            f"/projects/{project_id}/plannings/{planning_id}/validate", headers=headers
        )
        assert validate_response.status_code == 200
        reference_response = client.post(
            f"/projects/{project_id}/plannings/{planning_id}/reference", headers=headers
        )
        assert reference_response.status_code == 200

        clone_response = client.post(
            f"/projects/{project_id}/plannings",
            json={"source_planning_id": planning_id},
            headers=headers,
        )

        assert clone_response.status_code == 201
        cloned_tasks = {task["uid"]: task for task in clone_response.json()["tasks"]}
        assert cloned_tasks[100]["parent_uid"] is None
        assert cloned_tasks[2]["parent_uid"] == 100
        assert cloned_tasks[1]["parent_uid"] == 2


def test_creates_first_planning_from_a_hierarchical_legacy_project() -> None:
    # Correctness coverage for #160: create_planning's `else` branch -- copying legacy
    # MsTask rows into a project's very first WfPlanning, when neither
    # source_planning_id nor project.displayed_planning_id is set -- was hardened with
    # the same two-phase parent_uid assignment as the sibling source_planning_id branch
    # fixed in #103 (same composite self-reference hazard onto (project_id, uid) in
    # principle, since MsTask carries the same parent/child shape).
    #
    # Unlike #103, this is *not* a proven pre-fix regression test: extensive attempts
    # (adversarial uid gaps, individual per-row flush + UPDATE backfill to control
    # physical insertion order, both on SQLite and PostgreSQL) never got the unordered
    # `SELECT ... WHERE project_id = X` to actually return this table's rows
    # child-before-parent through create_planning's own query -- MsTask's query planning
    # behaves differently from WfPlanningTaskSnapshot's here, and it did not reproduce
    # the 409 either before or after the fix. The change is kept as defensive
    # hardening for symmetry and because MsTask.parent_uid could in principle be
    # populated out of order by future code (generate_planning_structure, the only
    # current writer, happens to always assign parent uids before children). This test
    # verifies the hierarchy still clones correctly, not that it previously failed.
    with TestClient(app) as client:
        headers = _auth_headers(client, "projects.legacy-hierarchical-clone@example.com")
        owner_id = _current_user_id(client, headers)

        session_factory = get_session_factory()
        with session_factory() as session:
            project = MsProject(
                owner_id=owner_id,
                source_version=2016,
                save_version_out=16,
                name="Legacy hierarchical clone source",
                schedule_from_start=True,
                start_date=datetime(2026, 1, 5, tzinfo=UTC),
                finish_date=datetime(2026, 1, 20, tzinfo=UTC),
                minutes_per_day=480,
                minutes_per_week=2400,
                days_per_month=20,
            )
            session.add(project)
            session.flush()
            project_id = project.id

            session.add(
                MsTask(
                    project_id=project_id,
                    uid=100,
                    parent_uid=None,
                    name="Root",
                    task_type=1,
                    outline_number="1",
                    outline_level=1,
                    is_summary=True,
                    is_milestone=False,
                )
            )
            session.add(
                MsTask(
                    project_id=project_id,
                    uid=2,
                    parent_uid=100,
                    name="Child",
                    task_type=0,
                    outline_number="1.1",
                    outline_level=2,
                    is_summary=False,
                    is_milestone=False,
                )
            )
            session.add(
                MsTask(
                    project_id=project_id,
                    uid=1,
                    parent_uid=2,
                    name="Grandchild",
                    task_type=0,
                    outline_number="1.1.1",
                    outline_level=3,
                    is_summary=False,
                    is_milestone=False,
                )
            )
            session.commit()

        # No source_planning_id and no displayed_planning_id yet: exercises
        # create_planning's `else` branch.
        clone_response = client.post(f"/projects/{project_id}/plannings", json={}, headers=headers)

        assert clone_response.status_code == 201
        cloned_tasks = {task["uid"]: task for task in clone_response.json()["tasks"]}
        assert cloned_tasks[100]["parent_uid"] is None
        assert cloned_tasks[2]["parent_uid"] == 100
        assert cloned_tasks[1]["parent_uid"] == 2


def test_creates_first_planning_from_legacy_project_preserves_task_enrichment_notes() -> None:
    # Regression for #178: create_planning's `else` branch (first clone of a legacy
    # project's MsTask rows, when neither source_planning_id nor
    # project.displayed_planning_id is set) built each cloned WfPlanningTaskSnapshot
    # without carrying over the task's notes, unlike the sibling source_planning_id
    # branch and reopen_planning_structure. MsTask itself has no `notes` column --
    # legacy task notes live in WfTaskEnrichment, keyed by (project_id, task_uid), the
    # same table update_task_description falls back to while no WfPlanning exists yet
    # (see tasks.py). This covers both a task with an enrichment row (its description
    # must land in WfPlanningTaskSnapshot.notes) and a task with none (must clone with
    # notes=None, not raise).
    with TestClient(app) as client:
        headers = _auth_headers(client, "projects.legacy-enrichment-clone@example.com")
        owner_id = _current_user_id(client, headers)

        session_factory = get_session_factory()
        with session_factory() as session:
            project = MsProject(
                owner_id=owner_id,
                source_version=2016,
                save_version_out=16,
                name="Legacy enrichment clone source",
                schedule_from_start=True,
                start_date=datetime(2026, 1, 5, tzinfo=UTC),
                finish_date=datetime(2026, 1, 20, tzinfo=UTC),
                minutes_per_day=480,
                minutes_per_week=2400,
                days_per_month=20,
            )
            session.add(project)
            session.flush()
            project_id = project.id

            session.add(
                MsTask(
                    project_id=project_id,
                    uid=1,
                    parent_uid=None,
                    name="Annotated task",
                    task_type=1,
                    outline_number="1",
                    outline_level=1,
                    is_summary=False,
                    is_milestone=False,
                )
            )
            session.add(
                MsTask(
                    project_id=project_id,
                    uid=2,
                    parent_uid=None,
                    name="Bare task",
                    task_type=1,
                    outline_number="2",
                    outline_level=1,
                    is_summary=False,
                    is_milestone=False,
                )
            )
            now = datetime.now(UTC)
            session.add(
                WfTaskEnrichment(
                    project_id=project_id,
                    task_uid=1,
                    description="Legacy notes must survive the first clone",
                    created_at=now,
                    updated_at=now,
                )
            )
            session.commit()

        # No source_planning_id and no displayed_planning_id yet: exercises
        # create_planning's `else` branch.
        clone_response = client.post(f"/projects/{project_id}/plannings", json={}, headers=headers)

        assert clone_response.status_code == 201
        cloned_tasks = {task["uid"]: task for task in clone_response.json()["tasks"]}
        assert cloned_tasks[1]["description"] == "Legacy notes must survive the first clone"
        assert cloned_tasks[2]["description"] is None


def test_projects_are_isolated_by_owner() -> None:
    with TestClient(app) as client:
        owner_headers = _auth_headers(client)
        owner_id = _current_user_id(client, owner_headers)
        project_id, _ = _seed_projects_and_tasks(owner_id)

        other_headers = _auth_headers(client, "projects.other@example.com")
        list_response: Response = client.get("/projects", headers=other_headers)
        assert list_response.status_code == 200
        other_projects = cast(list[dict[str, Any]], list_response.json()["items"])
        assert all(item["id"] != project_id for item in other_projects)

        for path in (
            f"/projects/{project_id}",
            f"/projects/{project_id}/tasks",
            f"/projects/{project_id}/export.xml",
        ):
            response: Response = client.get(path, headers=other_headers)
            assert response.status_code == 404

        update_response: Response = client.patch(
            f"/projects/{project_id}",
            json={"name": "Unauthorized"},
            headers=other_headers,
        )
        assert update_response.status_code == 404

        delete_response: Response = client.delete(f"/projects/{project_id}", headers=other_headers)
        assert delete_response.status_code == 404


def test_user_can_create_manual_project() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        response: Response = client.post(
            "/projects",
            json={"name": "Projet manuel", "currency_code": "eur"},
            headers=headers,
        )
        assert response.status_code == 201
        payload = cast(dict[str, Any], response.json())
        assert payload["name"] == "Projet manuel"
        assert payload["currency_code"] == "EUR"


def test_project_pagination_reports_total_and_task_listing_is_never_truncated() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        owner_id = _current_user_id(client, headers)
        first_project_id, expected_tasks = _seed_projects_and_tasks(owner_id)
        second_project_id, _ = _seed_projects_and_tasks(owner_id)

        first_page = client.get("/projects?limit=1&offset=0", headers=headers)
        second_page = client.get("/projects?limit=1&offset=1", headers=headers)
        assert first_page.status_code == 200
        assert second_page.status_code == 200
        first_body = cast(dict[str, Any], first_page.json())
        second_body = cast(dict[str, Any], second_page.json())
        assert first_body["limit"] == 1
        assert first_body["offset"] == 0
        assert second_body["offset"] == 1
        assert first_body["total"] == second_body["total"] == 2
        first_projects = cast(list[dict[str, Any]], first_body["items"])
        second_projects = cast(list[dict[str, Any]], second_body["items"])
        assert len(first_projects) == 1
        assert len(second_projects) == 1
        assert first_projects[0]["id"] != second_projects[0]["id"]

        without_limit = client.get("/projects", headers=headers)
        assert without_limit.status_code == 200
        without_limit_body = cast(dict[str, Any], without_limit.json())
        assert without_limit_body["limit"] is None
        assert without_limit_body["total"] == len(without_limit_body["items"]) == 2

        # /projects/{project_id}/tasks is intentionally not paginated (EPIC E7,
        # issue #115): it feeds the planning editor, which needs the whole task
        # tree, so any stray limit/offset a caller sends must be ignored rather
        # than silently truncating the response.
        task_page = client.get(
            f"/projects/{first_project_id}/tasks?limit=1&offset=1",
            headers=headers,
        )
        assert task_page.status_code == 200
        task_body = cast(dict[str, Any], task_page.json())
        assert task_body["limit"] is None
        assert task_body["offset"] == 0
        tasks = cast(list[dict[str, Any]], task_body["items"])
        assert len(tasks) == expected_tasks == task_body["total"]

        assert first_project_id != second_project_id


def test_list_projects_sort_and_rejects_unknown_sort_column() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client, "projects.sort@example.com")
        owner_id = _current_user_id(client, headers)
        _seed_projects_and_tasks(owner_id)
        created = client.post("/projects", json={"name": "Aardvark project"}, headers=headers)
        assert created.status_code == 201

        ascending = client.get("/projects?sort=name", headers=headers)
        assert ascending.status_code == 200
        names = cast(list[str], [item["name"] for item in ascending.json()["items"]])
        assert names == sorted(names)
        assert names[0] == "Aardvark project"

        invalid_sort = client.get("/projects?sort=unknown_column", headers=headers)
        assert invalid_sort.status_code == 400


def test_list_projects_reports_total_beyond_legacy_default_limit_of_50() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client, "projects.many@example.com")
        for index in range(55):
            created = client.post(
                "/projects", json={"name": f"Bulk project {index:03d}"}, headers=headers
            )
            assert created.status_code == 201

        # This is exactly the bug issue #115 fixes: GET /projects used to
        # default to limit=50 with no `total` in the response, silently
        # dropping anything past the 50th project with no signal to the
        # caller that more rows existed.
        response = client.get("/projects", headers=headers)
        assert response.status_code == 200
        body = cast(dict[str, Any], response.json())
        assert body["limit"] is None
        assert len(body["items"]) == 55
        assert body["total"] == 55

        capped = client.get("/projects?limit=10", headers=headers)
        assert capped.status_code == 200
        capped_body = cast(dict[str, Any], capped.json())
        assert capped_body["limit"] == 10
        assert len(capped_body["items"]) == 10
        assert capped_body["total"] == 55


def test_list_project_estimates_pagination_sort_and_validation() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))

        notes = ["Chiffrage initial detaille", "Autre devis"]
        for note in notes:
            created = client.post(
                f"/projects/{project_id}/estimates",
                json={"kind": "initial", "currency_code": "EUR", "note": note},
                headers=headers,
            )
            assert created.status_code == 201

        listed = client.get(f"/projects/{project_id}/estimates", headers=headers)
        assert listed.status_code == 200
        body = cast(dict[str, Any], listed.json())
        assert body["limit"] is None
        assert body["total"] == len(body["items"]) == 2

        descending = client.get(
            f"/projects/{project_id}/estimates?sort=-version_number", headers=headers
        )
        assert descending.status_code == 200
        descending_versions = [item["version_number"] for item in descending.json()["items"]]
        assert descending_versions == [2, 1]

        invalid_sort = client.get(
            f"/projects/{project_id}/estimates?sort=unknown_column", headers=headers
        )
        assert invalid_sort.status_code == 400

        searched = client.get(f"/projects/{project_id}/estimates?q=detaille", headers=headers)
        assert searched.status_code == 200
        assert [item["note"] for item in searched.json()["items"]] == ["Chiffrage initial detaille"]


def test_list_estimate_task_rows_pagination_sort_and_validation() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))
        estimate_id = cast(
            int,
            client.post(
                f"/projects/{project_id}/estimates",
                json={"kind": "initial", "currency_code": "EUR"},
                headers=headers,
            ).json()["id"],
        )

        listed = client.get(
            f"/projects/{project_id}/estimates/{estimate_id}/task-rows", headers=headers
        )
        assert listed.status_code == 200
        body = cast(dict[str, Any], listed.json())
        assert body["limit"] is None
        assert body["total"] == len(body["items"]) == 2

        descending = client.get(
            f"/projects/{project_id}/estimates/{estimate_id}/task-rows?sort=-task_name",
            headers=headers,
        )
        assert descending.status_code == 200
        names = [row["task_name"] for row in descending.json()["items"]]
        assert names == ["Task Two", "Task One"]

        invalid_sort = client.get(
            f"/projects/{project_id}/estimates/{estimate_id}/task-rows?sort=unknown_column",
            headers=headers,
        )
        assert invalid_sort.status_code == 400


def test_list_estimate_cost_lines_pagination_sort_and_validation() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))
        estimate_id = cast(
            int,
            client.post(
                f"/projects/{project_id}/estimates",
                json={"kind": "initial", "currency_code": "EUR"},
                headers=headers,
            ).json()["id"],
        )

        session_factory = get_session_factory()
        with session_factory() as session:
            cost_type = CostType(code=f"MAT-{uuid4().hex[:8]}", name="Materiel")
            session.add(cost_type)
            session.flush()
            cost_category = CostCategory(
                cost_type_id=cost_type.id,
                accounting_code=f"MATCAT-{uuid4().hex[:8]}",
                name="Materiel",
            )
            session.add(cost_category)
            session.commit()
            cost_category_id = cost_category.id

        for label in ("Bravo", "Alpha"):
            created = client.post(
                f"/projects/{project_id}/estimates/{estimate_id}/cost-lines",
                json={
                    "cost_category_id": cost_category_id,
                    "label": label,
                    "quantity": 1,
                    "unit_cost": 10,
                },
                headers=headers,
            )
            assert created.status_code == 201

        listed = client.get(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines", headers=headers
        )
        assert listed.status_code == 200
        body = cast(dict[str, Any], listed.json())
        assert body["limit"] is None
        assert body["total"] == len(body["items"]) == 2

        sorted_by_label = client.get(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines?sort=label",
            headers=headers,
        )
        assert sorted_by_label.status_code == 200
        labels = [line["label"] for line in sorted_by_label.json()["items"]]
        assert labels == ["Alpha", "Bravo"]

        invalid_sort = client.get(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines?sort=unknown_column",
            headers=headers,
        )
        assert invalid_sort.status_code == 400


def test_list_plannings_pagination_sort_and_validation() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))

        notes = ["Brouillon de reference", "Autre version"]
        for note in notes:
            created = client.post(
                f"/projects/{project_id}/plannings", json={"note": note}, headers=headers
            )
            assert created.status_code == 201

        listed = client.get(f"/projects/{project_id}/plannings", headers=headers)
        assert listed.status_code == 200
        body = cast(dict[str, Any], listed.json())
        assert body["limit"] is None
        assert body["total"] == len(body["items"]) == 2

        descending = client.get(
            f"/projects/{project_id}/plannings?sort=-version_number", headers=headers
        )
        assert descending.status_code == 200
        versions = [item["version_number"] for item in descending.json()["items"]]
        assert versions == [2, 1]

        invalid_sort = client.get(
            f"/projects/{project_id}/plannings?sort=unknown_column", headers=headers
        )
        assert invalid_sort.status_code == 400

        searched = client.get(f"/projects/{project_id}/plannings?q=reference", headers=headers)
        assert searched.status_code == 200
        assert [item["note"] for item in searched.json()["items"]] == ["Brouillon de reference"]


def test_list_task_role_assignments_pagination_sort_and_validation() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))
        labor_role_id, _ = _seed_roles()

        session_factory = get_session_factory()
        with session_factory() as session:
            labor_role = session.query(ResourceRole).filter(ResourceRole.id == labor_role_id).one()
            second_role = ResourceRole(
                node_id=labor_role.node_id,
                cost_category_id=labor_role.cost_category_id,
                name="Analyste",
            )
            session.add(second_role)
            session.commit()
            second_role_id = second_role.id

        first_assignment = client.post(
            f"/projects/{project_id}/tasks/1001/role-assignments",
            json={"role_id": labor_role_id, "quantity": "1", "hours": "1"},
            headers=headers,
        )
        assert first_assignment.status_code == 201
        second_assignment = client.post(
            f"/projects/{project_id}/tasks/1001/role-assignments",
            json={"role_id": second_role_id, "quantity": "1", "hours": "1"},
            headers=headers,
        )
        assert second_assignment.status_code == 201

        listed = client.get(f"/projects/{project_id}/tasks/1001/role-assignments", headers=headers)
        assert listed.status_code == 200
        body = cast(dict[str, Any], listed.json())
        assert body["limit"] is None
        assert body["total"] == len(body["items"]) == 2
        # Default sort is role_name ascending: "Analyste" precedes "Développeur".
        assert [item["role_name"] for item in body["items"]] == ["Analyste", "Développeur"]

        descending = client.get(
            f"/projects/{project_id}/tasks/1001/role-assignments?sort=-role_name",
            headers=headers,
        )
        assert descending.status_code == 200
        assert [item["role_name"] for item in descending.json()["items"]] == [
            "Développeur",
            "Analyste",
        ]

        invalid_sort = client.get(
            f"/projects/{project_id}/tasks/1001/role-assignments?sort=unknown_column",
            headers=headers,
        )
        assert invalid_sort.status_code == 400

        searched = client.get(
            f"/projects/{project_id}/tasks/1001/role-assignments?q=analyste",
            headers=headers,
        )
        assert searched.status_code == 200
        assert [item["role_name"] for item in searched.json()["items"]] == ["Analyste"]


def test_project_estimate_snapshots_tasks_and_validates() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))

        create_response = client.post(
            f"/projects/{project_id}/estimates",
            json={"kind": "initial", "currency_code": "EUR", "note": "Chiffrage initial"},
            headers=headers,
        )
        assert create_response.status_code == 201
        estimate = cast(dict[str, Any], create_response.json())
        estimate_id = cast(int, estimate["id"])
        assert estimate["version_number"] == 1
        assert estimate["status"] == "draft"

        rows_response = client.get(
            f"/projects/{project_id}/estimates/{estimate_id}/task-rows",
            headers=headers,
        )
        assert rows_response.status_code == 200
        rows = cast(list[dict[str, Any]], rows_response.json()["items"])
        assert [row["task_name"] for row in rows] == ["Task One", "Task Two"]
        assert [row["position"] for row in rows] == [1, 2]

        validate_response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/validate",
            headers=headers,
        )
        assert validate_response.status_code == 200
        assert validate_response.json()["status"] == "validated"
        assert validate_response.json()["validated_at"] is not None
        # Issue #65 (E6-04): neither task got a role assignment nor a cost line, so
        # both are surfaced as non-blocking warnings -- validation still succeeded.
        assert {w["task_uid"] for w in validate_response.json()["warnings"]} == {1001, 1002}

        validate_again_response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/validate",
            headers=headers,
        )
        assert validate_again_response.status_code == 409


def test_validate_estimate_warns_about_uncovered_tasks_only() -> None:
    """Issue #65 (E6-04): the validation warning only flags "real" tasks (excludes
    summaries/milestones) with neither a TaskRoleAssignment nor an EstimateCostLine
    of this estimate referencing them -- and never blocks validation."""
    with TestClient(app) as client:
        headers = _auth_headers(client, "projects.validation-warnings@example.com")
        owner_id = _current_user_id(client, headers)
        project_id, _ = _seed_projects_and_tasks(owner_id)
        labor_role_id, supply_role_id = _seed_roles()

        session_factory = get_session_factory()
        with session_factory() as session:
            summary_task = MsTask(
                project_id=project_id,
                uid=1003,
                name="Summary task",
                is_summary=True,
                is_milestone=False,
            )
            milestone_task = MsTask(
                project_id=project_id,
                uid=1004,
                name="Milestone task",
                is_summary=False,
                is_milestone=True,
            )
            uncovered_task = MsTask(
                project_id=project_id,
                uid=1005,
                name="Forgotten task",
                is_summary=False,
                is_milestone=False,
            )
            session.add_all([summary_task, milestone_task, uncovered_task])
            session.commit()

        tasks_response = client.get(f"/projects/{project_id}/tasks", headers=headers)
        assert tasks_response.status_code == 200
        tasks_by_uid = {
            task["uid"]: task["id"]
            for task in cast(list[dict[str, Any]], tasks_response.json()["items"])
        }

        estimate_response = client.post(
            f"/projects/{project_id}/estimates",
            json={"kind": "initial", "currency_code": "EUR"},
            headers=headers,
        )
        assert estimate_response.status_code == 201
        estimate_id = cast(int, estimate_response.json()["id"])

        # Task 1001 is covered via a TaskRoleAssignment.
        assignment_response = client.post(
            f"/projects/{project_id}/tasks/1001/role-assignments",
            json={"role_id": labor_role_id, "quantity": "1", "hours": "1"},
            headers=headers,
        )
        assert assignment_response.status_code == 201

        # Task 1002 is covered via an EstimateCostLine.task_id.
        cost_line_response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines",
            json={
                "task_id": tasks_by_uid[1002],
                "cost_category_id": _cost_category_id_for_role(supply_role_id),
                "label": "Câble dédié",
                "quantity": 1,
                "unit_cost": 10,
            },
            headers=headers,
        )
        assert cost_line_response.status_code == 201

        # A cost line with no task_id at all must not count as covering any task.
        unattached_cost_line_response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines",
            json={
                "cost_category_id": _cost_category_id_for_role(supply_role_id),
                "label": "Frais generaux",
                "quantity": 1,
                "unit_cost": 5,
            },
            headers=headers,
        )
        assert unattached_cost_line_response.status_code == 201

        validate_response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/validate",
            headers=headers,
        )
        assert validate_response.status_code == 200
        body = cast(dict[str, Any], validate_response.json())
        assert body["status"] == "validated"
        assert body["warnings"] == [{"task_uid": 1005, "task_name": "Forgotten task"}]


def test_validate_estimate_reports_no_warnings_when_all_tasks_are_covered() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client, "projects.validation-no-warnings@example.com")
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))
        labor_role_id, supply_role_id = _seed_roles()

        tasks_response = client.get(f"/projects/{project_id}/tasks", headers=headers)
        assert tasks_response.status_code == 200
        tasks_by_uid = {
            task["uid"]: task["id"]
            for task in cast(list[dict[str, Any]], tasks_response.json()["items"])
        }

        estimate_response = client.post(
            f"/projects/{project_id}/estimates",
            json={"kind": "initial", "currency_code": "EUR"},
            headers=headers,
        )
        assert estimate_response.status_code == 201
        estimate_id = cast(int, estimate_response.json()["id"])

        assert (
            client.post(
                f"/projects/{project_id}/tasks/1001/role-assignments",
                json={"role_id": labor_role_id, "quantity": "1", "hours": "1"},
                headers=headers,
            ).status_code
            == 201
        )
        assert (
            client.post(
                f"/projects/{project_id}/estimates/{estimate_id}/cost-lines",
                json={
                    "task_id": tasks_by_uid[1002],
                    "cost_category_id": _cost_category_id_for_role(supply_role_id),
                    "label": "Câble dédié",
                    "quantity": 1,
                    "unit_cost": 10,
                },
                headers=headers,
            ).status_code
            == 201
        )

        validate_response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/validate",
            headers=headers,
        )
        assert validate_response.status_code == 200
        assert validate_response.json()["warnings"] == []


def test_project_estimate_can_snapshot_planning_without_legacy_tasks() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client, "projects.snapshot-only@example.com")
        project_response = client.post(
            "/projects",
            json={"name": "Snapshot-only project"},
            headers=headers,
        )
        assert project_response.status_code == 201
        project_payload = cast(dict[str, Any], project_response.json())
        project_id = cast(int, project_payload["id"])

        session_factory = get_session_factory()
        with session_factory() as session:
            planning = WfPlanning(project_id=project_id, version_number=1, status="draft")
            session.add(planning)
            session.flush()
            session.add(
                WfPlanningTaskSnapshot(
                    planning_id=planning.id,
                    uid=9001,
                    name="Snapshot task",
                    position=1,
                    is_summary=False,
                    is_milestone=False,
                )
            )
            project = session.query(MsProject).filter(MsProject.id == project_id).one()
            project.displayed_planning_id = planning.id
            session.commit()

        estimate_response = client.post(
            f"/projects/{project_id}/estimates",
            json={"kind": "initial", "currency_code": "EUR"},
            headers=headers,
        )
        assert estimate_response.status_code == 201
        estimate_id = cast(int, estimate_response.json()["id"])

        rows_response = client.get(
            f"/projects/{project_id}/estimates/{estimate_id}/task-rows",
            headers=headers,
        )
        assert rows_response.status_code == 200
        assert rows_response.json()["items"][0]["task_id"] is None
        assert rows_response.json()["items"][0]["task_name"] == "Snapshot task"


def test_planning_lifecycle_snapshots_reference_and_display_selection() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))

        create_response = client.post(
            f"/projects/{project_id}/plannings",
            json={"note": "Version initiale"},
            headers=headers,
        )
        assert create_response.status_code == 201
        draft = cast(dict[str, Any], create_response.json())
        planning_id = cast(int, draft["id"])
        assert draft["status"] == "draft"
        assert [task["uid"] for task in draft["tasks"]] == [1001, 1002]

        invalid_reference = client.post(
            f"/projects/{project_id}/plannings/{planning_id}/reference",
            headers=headers,
        )
        assert invalid_reference.status_code == 409

        validate_response = client.post(
            f"/projects/{project_id}/plannings/{planning_id}/validate",
            headers=headers,
        )
        assert validate_response.status_code == 200

        first_reference = client.post(
            f"/projects/{project_id}/plannings/{planning_id}/reference",
            headers=headers,
        )
        assert first_reference.status_code == 200
        assert first_reference.json()["planning_reference_id"] == planning_id

        second_response = client.post(
            f"/projects/{project_id}/plannings",
            json={"source_planning_id": planning_id},
            headers=headers,
        )
        assert second_response.status_code == 201
        second_id = second_response.json()["id"]
        assert second_response.json()["tasks"][0]["name"] == "Task One"
        assert (
            client.post(
                f"/projects/{project_id}/plannings/{second_id}/validate", headers=headers
            ).status_code
            == 200
        )
        assert (
            client.post(
                f"/projects/{project_id}/plannings/{second_id}/reference", headers=headers
            ).status_code
            == 200
        )

        with get_session_factory()() as session:
            first = session.get(WfPlanning, planning_id)
            assert first is not None
            assert first.status == "superseded"

        displayed = client.post(
            f"/projects/{project_id}/plannings/{planning_id}/display", headers=headers
        )
        assert displayed.status_code == 200
        tasks = client.get(f"/projects/{project_id}/tasks", headers=headers)
        assert tasks.status_code == 200
        assert tasks.json()["items"][0]["name"] == "Task One"


def test_project_status_requires_references_and_excludes_archived_by_default() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))

        blocked = client.patch(
            f"/projects/{project_id}/status",
            json={"status": "en_cours"},
            headers=headers,
        )
        assert blocked.status_code == 409

        completed = client.patch(
            f"/projects/{project_id}/status",
            json={"status": "termine"},
            headers=headers,
        )
        assert completed.status_code == 409
        lost = client.patch(
            f"/projects/{project_id}/status",
            json={"status": "perdu"},
            headers=headers,
        )
        assert lost.status_code == 200
        assert lost.json()["status"] == "perdu"
        active_projects = cast(
            list[dict[str, Any]], client.get("/projects", headers=headers).json()["items"]
        )
        assert all(item["id"] != project_id for item in active_projects)
        archived_projects = cast(
            list[dict[str, Any]],
            client.get("/projects?include_archived=true", headers=headers).json()["items"],
        )
        assert any(item["id"] == project_id for item in archived_projects)


def test_project_can_enter_in_progress_after_both_references_are_set() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))

        planning = client.post(f"/projects/{project_id}/plannings", json={}, headers=headers)
        assert planning.status_code == 201
        planning_id = planning.json()["id"]
        assert (
            client.post(
                f"/projects/{project_id}/plannings/{planning_id}/validate", headers=headers
            ).status_code
            == 200
        )
        assert (
            client.post(
                f"/projects/{project_id}/plannings/{planning_id}/reference", headers=headers
            ).status_code
            == 200
        )

        estimate = client.post(
            f"/projects/{project_id}/estimates",
            json={"kind": "initial", "currency_code": "EUR"},
            headers=headers,
        )
        assert estimate.status_code == 201
        estimate_id = estimate.json()["id"]
        assert (
            client.post(
                f"/projects/{project_id}/estimates/{estimate_id}/validate", headers=headers
            ).status_code
            == 200
        )
        assert (
            client.post(
                f"/projects/{project_id}/estimates/{estimate_id}/reference", headers=headers
            ).status_code
            == 200
        )

        assert (
            client.patch(
                f"/projects/{project_id}/status",
                json={"status": "initialise"},
                headers=headers,
            ).status_code
            == 200
        )
        assert (
            client.patch(
                f"/projects/{project_id}/status",
                json={"status": "en_reponse_appel_offre"},
                headers=headers,
            ).status_code
            == 200
        )

        response = client.patch(
            f"/projects/{project_id}/status",
            json={"status": "en_cours"},
            headers=headers,
        )
        assert response.status_code == 200
        assert response.json()["status"] == "en_cours"


def test_en_cours_transition_rejects_project_initialised_without_structure_via_skip() -> None:
    """Regression test for the second review round on #130.

    Since the ``cree -> initialise`` structure gate was removed, a project
    can reach ``initialise`` with zero tasks by skipping the poste/lot/
    livrable step. Before this fix, the remaining ``en_cours`` structure
    check was dead code: ``project.status in {"cree", "initialise"}`` is
    always false by the time the transition is evaluated, since the only
    legal source status for ``en_cours`` is ``en_reponse_appel_offre``. A
    project could therefore reach ``en_cours`` without ever having a
    planning structure or any task.
    """
    with TestClient(app) as client:
        headers = _auth_headers(client)
        create = client.post("/projects", json={"name": "Skip then en_cours"}, headers=headers)
        assert create.status_code == 201
        project_id = create.json()["id"]

        skip = client.post(f"/projects/{project_id}/planning-structure/skip", headers=headers)
        assert skip.status_code == 200
        assert skip.json()["status"] == "initialise"
        planning_id = skip.json()["displayed_planning_id"]
        assert planning_id is not None

        assert (
            client.post(
                f"/projects/{project_id}/plannings/{planning_id}/validate", headers=headers
            ).status_code
            == 200
        )
        assert (
            client.post(
                f"/projects/{project_id}/plannings/{planning_id}/reference", headers=headers
            ).status_code
            == 200
        )
        assert client.get(f"/projects/{project_id}/tasks", headers=headers).json()["items"] == []

        estimate = client.post(
            f"/projects/{project_id}/estimates",
            json={"kind": "initial", "currency_code": "EUR"},
            headers=headers,
        )
        assert estimate.status_code == 201
        estimate_id = estimate.json()["id"]
        assert (
            client.post(
                f"/projects/{project_id}/estimates/{estimate_id}/validate", headers=headers
            ).status_code
            == 200
        )
        assert (
            client.post(
                f"/projects/{project_id}/estimates/{estimate_id}/reference", headers=headers
            ).status_code
            == 200
        )

        assert (
            client.patch(
                f"/projects/{project_id}/status",
                json={"status": "en_reponse_appel_offre"},
                headers=headers,
            ).status_code
            == 200
        )

        response = client.patch(
            f"/projects/{project_id}/status",
            json={"status": "en_cours"},
            headers=headers,
        )
        assert response.status_code == 409
        assert response.json()["detail"] == {"code": "GENERIC_ERROR"}


def test_en_cours_transition_rejects_structure_in_unrelated_draft_not_referenced() -> None:
    """Regression test for the third review round on #130.

    ``_project_has_structure`` searched every planning version of the
    project, so a project could satisfy the ``en_cours`` structure
    requirement via an unrelated/historical draft while the actually
    referenced planning was empty. Build exactly that trap: an empty
    referenced planning (v1) coexists with a second, unreferenced draft
    (v2) that does have a task, and the transition must still be rejected.
    """
    with TestClient(app) as client:
        headers = _auth_headers(client)
        create = client.post(
            "/projects", json={"name": "Structure in wrong draft"}, headers=headers
        )
        assert create.status_code == 201
        project_id = create.json()["id"]

        skip = client.post(f"/projects/{project_id}/planning-structure/skip", headers=headers)
        assert skip.status_code == 200
        assert skip.json()["status"] == "initialise"
        v1_id = skip.json()["displayed_planning_id"]
        assert v1_id is not None

        v2 = client.post(f"/projects/{project_id}/plannings", json={}, headers=headers)
        assert v2.status_code == 201
        v2_id = v2.json()["id"]
        task_created = client.post(
            f"/projects/{project_id}/plannings/{v2_id}/tasks",
            json={"name": "Structure elsewhere", "expected_revision": 0},
            headers=headers,
        )
        assert task_created.status_code == 200

        assert (
            client.post(
                f"/projects/{project_id}/plannings/{v1_id}/validate", headers=headers
            ).status_code
            == 200
        )
        assert (
            client.post(
                f"/projects/{project_id}/plannings/{v1_id}/reference", headers=headers
            ).status_code
            == 200
        )
        assert (
            client.get(f"/projects/{project_id}", headers=headers).json()["planning_reference_id"]
            == v1_id
        )

        estimate = client.post(
            f"/projects/{project_id}/estimates",
            json={"kind": "initial", "currency_code": "EUR"},
            headers=headers,
        )
        assert estimate.status_code == 201
        estimate_id = estimate.json()["id"]
        assert (
            client.post(
                f"/projects/{project_id}/estimates/{estimate_id}/validate", headers=headers
            ).status_code
            == 200
        )
        assert (
            client.post(
                f"/projects/{project_id}/estimates/{estimate_id}/reference", headers=headers
            ).status_code
            == 200
        )

        assert (
            client.patch(
                f"/projects/{project_id}/status",
                json={"status": "en_reponse_appel_offre"},
                headers=headers,
            ).status_code
            == 200
        )

        response = client.patch(
            f"/projects/{project_id}/status",
            json={"status": "en_cours"},
            headers=headers,
        )
        assert response.status_code == 409
        assert response.json()["detail"] == {"code": "GENERIC_ERROR"}


def test_project_can_initialise_without_a_planning_structure() -> None:
    """Issue #130: the poste/lot/livrable skeleton is optional before ``initialise``."""
    with TestClient(app) as client:
        headers = _auth_headers(client)
        create = client.post("/projects", json={"name": "No skeleton"}, headers=headers)
        assert create.status_code == 201
        project_id = create.json()["id"]

        response = client.patch(
            f"/projects/{project_id}/status",
            json={"status": "initialise"},
            headers=headers,
        )

        assert response.status_code == 200
        assert response.json()["status"] == "initialise"


def test_estimate_cost_lines_support_non_labor_costs_and_draft_locking() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))
        labor_role_id, supply_role_id = _seed_roles()

        session_factory = get_session_factory()
        with session_factory() as session:
            labor_category_id = (
                session.query(ResourceRole)
                .filter(ResourceRole.id == labor_role_id)
                .one()
                .cost_category_id
            )
            supply_category_id = (
                session.query(ResourceRole)
                .filter(ResourceRole.id == supply_role_id)
                .one()
                .cost_category_id
            )
            fee_type = CostType(code="FRAIS", name="Frais", kind="other")
            work_unit_type = CostType(code="UO", name="Unité d'oeuvre", kind="other")
            session.add_all([fee_type, work_unit_type])
            session.flush()
            fee_category = CostCategory(
                cost_type_id=fee_type.id,
                accounting_code="FR-TRAVEL",
                category_code="FRAIS",
                name="Déplacement",
            )
            work_unit_category = CostCategory(
                cost_type_id=work_unit_type.id,
                accounting_code="UO-TEST",
                category_code="UO",
                name="Essais",
            )
            session.add_all([fee_category, work_unit_category])
            session.commit()
            fee_category_id = fee_category.id
            work_unit_category_id = work_unit_category.id

        estimate_response = client.post(
            f"/projects/{project_id}/estimates",
            json={"kind": "initial", "currency_code": "EUR"},
            headers=headers,
        )
        assert estimate_response.status_code == 201
        estimate_id = cast(int, estimate_response.json()["id"])

        create_response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines",
            json={
                "task_id": 1,
                "cost_category_id": supply_category_id,
                "label": "Câble réseau",
                "quantity": "3",
                "unit_cost": "12.50",
                "supply_status": "ordered",
            },
            headers=headers,
        )
        assert create_response.status_code == 201
        cost_line = cast(dict[str, Any], create_response.json())
        cost_line_id = cast(int, cost_line["id"])
        assert cost_line["cost_type_code"] == "FOURNITURE"
        assert cost_line["accounting_code"] == "FO-CABLE"
        assert cost_line["category_code"] == "ACHAT"
        assert cost_line["purchase_cost"] == "37.50"
        assert cost_line["supply_status"] == "ordered"

        other_headers = _auth_headers(client, "cost-line.other@example.com")
        other_response = client.get(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines",
            headers=other_headers,
        )
        assert other_response.status_code == 404
        other_delete_response = client.delete(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines/{cost_line_id}",
            headers=other_headers,
        )
        assert other_delete_response.status_code == 404

        fee_response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines",
            json={
                "cost_category_id": fee_category_id,
                "label": "Déplacement",
                "quantity": "2",
                "unit_cost": "120",
            },
            headers=headers,
        )
        assert fee_response.status_code == 201
        assert fee_response.json()["cost_type_code"] == "FRAIS"
        assert fee_response.json()["task_id"] is None
        assert fee_response.json()["supply_status"] is None

        work_unit_response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines",
            json={
                "cost_category_id": work_unit_category_id,
                "label": "Essais laboratoire",
                "quantity": "4",
                "unit_cost": "80",
            },
            headers=headers,
        )
        assert work_unit_response.status_code == 201
        assert work_unit_response.json()["cost_type_code"] == "UO"

        status_for_fee_response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines",
            json={
                "cost_category_id": fee_category_id,
                "label": "Statut non valide",
                "quantity": "1",
                "unit_cost": "1",
                "supply_status": "ordered",
            },
            headers=headers,
        )
        assert status_for_fee_response.status_code == 400

        labor_response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines",
            json={
                "cost_category_id": labor_category_id,
                "label": "Non autorisé",
                "quantity": "1",
                "unit_cost": "1",
            },
            headers=headers,
        )
        assert labor_response.status_code == 400

        invalid_task_response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines",
            json={
                "task_id": 999999,
                "cost_category_id": supply_category_id,
                "label": "Tâche absente",
                "quantity": "1",
                "unit_cost": "1",
            },
            headers=headers,
        )
        assert invalid_task_response.status_code == 400

        invalid_status_response = client.patch(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines/{cost_line_id}",
            json={"supply_status": "received", "unit_cost": "15"},
            headers=headers,
        )
        assert invalid_status_response.status_code == 200
        assert invalid_status_response.json()["purchase_cost"] == "45.00"

        validate_response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/validate",
            headers=headers,
        )
        assert validate_response.status_code == 200

        locked_response = client.delete(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines/{cost_line_id}",
            headers=headers,
        )
        assert locked_response.status_code == 409


def test_estimate_cost_line_defaults_to_project_root_cost_code() -> None:
    """Issue #63 (E6-02): a non-labor cost line created without an explicit
    `cost_code_id` is attached to the project's active root cost code by default."""
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))
        _, supply_role_id = _seed_roles()
        root_id = _root_cost_code_id(client, headers, project_id)

        session_factory = get_session_factory()
        with session_factory() as session:
            supply_category_id = (
                session.query(ResourceRole)
                .filter(ResourceRole.id == supply_role_id)
                .one()
                .cost_category_id
            )

        estimate_response = client.post(
            f"/projects/{project_id}/estimates",
            json={"kind": "initial", "currency_code": "EUR"},
            headers=headers,
        )
        estimate_id = cast(int, estimate_response.json()["id"])

        create_response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines",
            json={
                "cost_category_id": supply_category_id,
                "label": "Câble réseau",
                "quantity": "1",
                "unit_cost": "1",
            },
            headers=headers,
        )
        assert create_response.status_code == 201
        assert create_response.json()["cost_code_id"] == root_id


def test_estimate_cost_line_accepts_explicit_sub_cost_code_and_rejects_foreign_one() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))
        _, supply_role_id = _seed_roles()
        root_id = _root_cost_code_id(client, headers, project_id)

        session_factory = get_session_factory()
        with session_factory() as session:
            supply_category_id = (
                session.query(ResourceRole)
                .filter(ResourceRole.id == supply_role_id)
                .one()
                .cost_category_id
            )

        sub_code_response = client.post(
            f"/projects/{project_id}/cost-codes",
            json={"code": "LOT-FOURN", "name": "Lot fourniture", "parent_id": root_id},
            headers=headers,
        )
        assert sub_code_response.status_code == 201
        sub_code_id = cast(int, sub_code_response.json()["id"])

        estimate_response = client.post(
            f"/projects/{project_id}/estimates",
            json={"kind": "initial", "currency_code": "EUR"},
            headers=headers,
        )
        estimate_id = cast(int, estimate_response.json()["id"])

        create_response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines",
            json={
                "cost_category_id": supply_category_id,
                "label": "Câble réseau",
                "quantity": "1",
                "unit_cost": "1",
                "cost_code_id": sub_code_id,
            },
            headers=headers,
        )
        assert create_response.status_code == 201
        cost_line_id = cast(int, create_response.json()["id"])
        assert create_response.json()["cost_code_id"] == sub_code_id

        other_project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))
        other_root_id = _root_cost_code_id(client, headers, other_project_id)

        rejected_create_response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines",
            json={
                "cost_category_id": supply_category_id,
                "label": "Câble refusé",
                "quantity": "1",
                "unit_cost": "1",
                "cost_code_id": other_root_id,
            },
            headers=headers,
        )
        assert rejected_create_response.status_code == 400

        rejected_update_response = client.patch(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines/{cost_line_id}",
            json={"cost_code_id": other_root_id},
            headers=headers,
        )
        assert rejected_update_response.status_code == 400

        accepted_update_response = client.patch(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines/{cost_line_id}",
            json={"cost_code_id": root_id},
            headers=headers,
        )
        assert accepted_update_response.status_code == 200
        assert accepted_update_response.json()["cost_code_id"] == root_id


def test_estimate_cost_line_explicit_null_cost_code_id_resolves_to_root_not_cleared() -> None:
    """Review finding on #63: `{"cost_code_id": null}` must resolve back to the
    project's root, exactly like an omitted field at create time -- never leave the
    line with no attachment at all."""
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))
        _, supply_role_id = _seed_roles()
        root_id = _root_cost_code_id(client, headers, project_id)

        session_factory = get_session_factory()
        with session_factory() as session:
            supply_category_id = (
                session.query(ResourceRole)
                .filter(ResourceRole.id == supply_role_id)
                .one()
                .cost_category_id
            )

        sub_code_response = client.post(
            f"/projects/{project_id}/cost-codes",
            json={"code": "LOT-NULL", "name": "Lot", "parent_id": root_id},
            headers=headers,
        )
        sub_code_id = cast(int, sub_code_response.json()["id"])

        estimate_id = cast(
            int,
            client.post(
                f"/projects/{project_id}/estimates",
                json={"kind": "initial", "currency_code": "EUR"},
                headers=headers,
            ).json()["id"],
        )

        create_response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines",
            json={
                "cost_category_id": supply_category_id,
                "label": "Câble",
                "quantity": "1",
                "unit_cost": "1",
                "cost_code_id": sub_code_id,
            },
            headers=headers,
        )
        cost_line_id = cast(int, create_response.json()["id"])
        assert create_response.json()["cost_code_id"] == sub_code_id

        null_update_response = client.patch(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines/{cost_line_id}",
            json={"cost_code_id": None},
            headers=headers,
        )
        assert null_update_response.status_code == 200
        assert null_update_response.json()["cost_code_id"] == root_id


def test_estimate_cost_line_rejects_a_deactivated_cost_code() -> None:
    """Review finding on #63: a deactivated cost code must never accept a new
    attachment, even though it may still be referenced by pre-existing lines."""
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))
        _, supply_role_id = _seed_roles()
        root_id = _root_cost_code_id(client, headers, project_id)

        session_factory = get_session_factory()
        with session_factory() as session:
            supply_category_id = (
                session.query(ResourceRole)
                .filter(ResourceRole.id == supply_role_id)
                .one()
                .cost_category_id
            )

        sub_code_response = client.post(
            f"/projects/{project_id}/cost-codes",
            json={"code": "LOT-INACTIVE", "name": "Lot", "parent_id": root_id},
            headers=headers,
        )
        sub_code_id = cast(int, sub_code_response.json()["id"])

        deactivate_response = client.delete(
            f"/projects/{project_id}/cost-codes/{sub_code_id}", headers=headers
        )
        assert deactivate_response.status_code == 204

        estimate_id = cast(
            int,
            client.post(
                f"/projects/{project_id}/estimates",
                json={"kind": "initial", "currency_code": "EUR"},
                headers=headers,
            ).json()["id"],
        )
        rejected_create_response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines",
            json={
                "cost_category_id": supply_category_id,
                "label": "Câble refusé",
                "quantity": "1",
                "unit_cost": "1",
                "cost_code_id": sub_code_id,
            },
            headers=headers,
        )
        assert rejected_create_response.status_code == 400


def test_estimate_cost_line_planned_date_is_independent_from_task_id() -> None:
    """Issue #66 (E6-05): `planned_date` and `task_id` are entirely independent --
    either, both, or neither may be set on a given cost line."""
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))
        _, supply_role_id = _seed_roles()

        session_factory = get_session_factory()
        with session_factory() as session:
            supply_category_id = (
                session.query(ResourceRole)
                .filter(ResourceRole.id == supply_role_id)
                .one()
                .cost_category_id
            )

        estimate_id = cast(
            int,
            client.post(
                f"/projects/{project_id}/estimates",
                json={"kind": "initial", "currency_code": "EUR"},
                headers=headers,
            ).json()["id"],
        )

        # planned_date without task_id.
        date_only_response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines",
            json={
                "cost_category_id": supply_category_id,
                "label": "Câble avec date",
                "quantity": "1",
                "unit_cost": "1",
                "planned_date": "2026-10-01T00:00:00Z",
            },
            headers=headers,
        )
        assert date_only_response.status_code == 201
        assert date_only_response.json()["task_id"] is None
        assert date_only_response.json()["planned_date"].startswith("2026-10-01T00:00:00")

        # task_id without planned_date.
        task_only_response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines",
            json={
                "task_id": 1,
                "cost_category_id": supply_category_id,
                "label": "Câble avec tâche",
                "quantity": "1",
                "unit_cost": "1",
            },
            headers=headers,
        )
        assert task_only_response.status_code == 201
        assert task_only_response.json()["task_id"] == 1
        assert task_only_response.json()["planned_date"] is None

        # Neither set.
        neither_response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines",
            json={
                "cost_category_id": supply_category_id,
                "label": "Câble sans date ni tâche",
                "quantity": "1",
                "unit_cost": "1",
            },
            headers=headers,
        )
        assert neither_response.status_code == 201
        assert neither_response.json()["task_id"] is None
        assert neither_response.json()["planned_date"] is None

        # Both set.
        both_response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines",
            json={
                "task_id": 1,
                "cost_category_id": supply_category_id,
                "label": "Câble avec date et tâche",
                "quantity": "1",
                "unit_cost": "1",
                "planned_date": "2026-11-15T08:30:00Z",
            },
            headers=headers,
        )
        assert both_response.status_code == 201
        assert both_response.json()["task_id"] == 1
        assert both_response.json()["planned_date"].startswith("2026-11-15T08:30:00")


def test_estimate_cost_line_planned_date_is_editable_independently_of_task_id() -> None:
    """Issue #66 (E6-05): `planned_date` is editable via PATCH independently of
    `task_id`, and an explicit `null` clears it -- consistent with how `task_id`
    already accepts an explicit `null` to detach a line from its task."""
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))
        _, supply_role_id = _seed_roles()

        session_factory = get_session_factory()
        with session_factory() as session:
            supply_category_id = (
                session.query(ResourceRole)
                .filter(ResourceRole.id == supply_role_id)
                .one()
                .cost_category_id
            )

        estimate_id = cast(
            int,
            client.post(
                f"/projects/{project_id}/estimates",
                json={"kind": "initial", "currency_code": "EUR"},
                headers=headers,
            ).json()["id"],
        )

        create_response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines",
            json={
                "task_id": 1,
                "cost_category_id": supply_category_id,
                "label": "Câble",
                "quantity": "1",
                "unit_cost": "1",
            },
            headers=headers,
        )
        assert create_response.status_code == 201
        cost_line_id = cast(int, create_response.json()["id"])
        assert create_response.json()["planned_date"] is None

        # Set planned_date without touching task_id.
        set_date_response = client.patch(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines/{cost_line_id}",
            json={"planned_date": "2026-12-01T00:00:00Z"},
            headers=headers,
        )
        assert set_date_response.status_code == 200
        assert set_date_response.json()["task_id"] == 1
        assert set_date_response.json()["planned_date"].startswith("2026-12-01T00:00:00")

        # Clear task_id without touching planned_date.
        clear_task_response = client.patch(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines/{cost_line_id}",
            json={"task_id": None},
            headers=headers,
        )
        assert clear_task_response.status_code == 200
        assert clear_task_response.json()["task_id"] is None
        assert clear_task_response.json()["planned_date"].startswith("2026-12-01T00:00:00")

        # Explicit null clears planned_date.
        clear_date_response = client.patch(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines/{cost_line_id}",
            json={"planned_date": None},
            headers=headers,
        )
        assert clear_date_response.status_code == 200
        assert clear_date_response.json()["planned_date"] is None


def test_estimate_line_snapshot_carries_source_cost_code_id() -> None:
    """Issue #63 (E6-02): calculate_estimate_lines copies `cost_code_id` from its
    source (TaskRoleAssignment for labor, EstimateCostLine for non-labor) onto the
    frozen EstimateLine snapshot at validation time, independently of
    `accounting_code`."""
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))
        labor_role_id, supply_role_id = _seed_roles()
        root_id = _root_cost_code_id(client, headers, project_id)

        session_factory = get_session_factory()
        with session_factory() as session:
            supply_category_id = (
                session.query(ResourceRole)
                .filter(ResourceRole.id == supply_role_id)
                .one()
                .cost_category_id
            )

        labor_sub_code_response = client.post(
            f"/projects/{project_id}/cost-codes",
            json={"code": "LOT-MO-SNAP", "name": "Lot MO", "parent_id": root_id},
            headers=headers,
        )
        assert labor_sub_code_response.status_code == 201
        labor_sub_code_id = cast(int, labor_sub_code_response.json()["id"])

        supply_sub_code_response = client.post(
            f"/projects/{project_id}/cost-codes",
            json={"code": "LOT-FOURN-SNAP", "name": "Lot fourniture", "parent_id": root_id},
            headers=headers,
        )
        assert supply_sub_code_response.status_code == 201
        supply_sub_code_id = cast(int, supply_sub_code_response.json()["id"])

        estimate_response = client.post(
            f"/projects/{project_id}/estimates",
            json={"kind": "initial", "currency_code": "EUR"},
            headers=headers,
        )
        estimate_id = cast(int, estimate_response.json()["id"])

        assignment_response = client.post(
            f"/projects/{project_id}/tasks/1001/role-assignments",
            json={
                "role_id": labor_role_id,
                "quantity": "1",
                "hours": "1",
                "cost_code_id": labor_sub_code_id,
            },
            headers=headers,
        )
        assert assignment_response.status_code == 201

        cost_line_response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/cost-lines",
            json={
                "cost_category_id": supply_category_id,
                "label": "Fourniture snapshot",
                "quantity": "1",
                "unit_cost": "1",
                "cost_code_id": supply_sub_code_id,
            },
            headers=headers,
        )
        assert cost_line_response.status_code == 201

        validate_response = client.post(
            f"/projects/{project_id}/estimates/{estimate_id}/validate",
            headers=headers,
        )
        assert validate_response.status_code == 200

        with session_factory() as session:
            from waterfall.models.resources import EstimateLine

            lines = (
                session.query(EstimateLine)
                .filter(EstimateLine.estimate_id == estimate_id)
                .order_by(EstimateLine.role_id.is_(None))
                .all()
            )
            cost_code_ids_by_role_presence = {
                line.role_id is not None: line.cost_code_id for line in lines
            }
            assert cost_code_ids_by_role_presence[True] == labor_sub_code_id
            assert cost_code_ids_by_role_presence[False] == supply_sub_code_id


def test_forecast_estimate_requires_same_project_reference() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))

        missing_reference_response = client.post(
            f"/projects/{project_id}/estimates",
            json={"kind": "forecast_remaining", "currency_code": "EUR"},
            headers=headers,
        )
        assert missing_reference_response.status_code == 400

        initial_response = client.post(
            f"/projects/{project_id}/estimates",
            json={"kind": "initial", "currency_code": "EUR"},
            headers=headers,
        )
        assert initial_response.status_code == 201
        initial_id = initial_response.json()["id"]

        forecast_response = client.post(
            f"/projects/{project_id}/estimates",
            json={
                "kind": "forecast_remaining",
                "currency_code": "EUR",
                "reference_estimate_id": initial_id,
            },
            headers=headers,
        )
        assert forecast_response.status_code == 201
        assert forecast_response.json()["reference_estimate_id"] == initial_id
        assert forecast_response.json()["version_number"] == 2

        other_headers = _auth_headers(client, "estimate.other@example.com")
        hidden_response = client.get(
            f"/projects/{project_id}/estimates/{initial_id}",
            headers=other_headers,
        )
        assert hidden_response.status_code == 404


def test_task_role_assignment_lifecycle_and_labor_validation() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))
        labor_role_id, supply_role_id = _seed_roles()

        rejected_response = client.post(
            f"/projects/{project_id}/tasks/1001/role-assignments",
            json={"role_id": supply_role_id, "quantity": "1", "hours": "7.4"},
            headers=headers,
        )
        assert rejected_response.status_code == 400

        create_response = client.post(
            f"/projects/{project_id}/tasks/1001/role-assignments",
            json={"role_id": labor_role_id, "quantity": "2", "hours": "7.4"},
            headers=headers,
        )
        assert create_response.status_code == 201
        assignment = cast(dict[str, Any], create_response.json())
        assignment_id = cast(int, assignment["id"])
        assert assignment["role_code"] == "Développeur"
        assert assignment["accounting_code"] == "MO-DEV"

        list_response = client.get(
            f"/projects/{project_id}/tasks/1001/role-assignments",
            headers=headers,
        )
        assert list_response.status_code == 200
        assignments = cast(list[dict[str, Any]], list_response.json()["items"])
        assert len(assignments) == 1

        update_response = client.patch(
            f"/projects/{project_id}/tasks/1001/role-assignments/{assignment_id}",
            json={"hours": "14.8"},
            headers=headers,
        )
        assert update_response.status_code == 200
        assert update_response.json()["hours"] == "14.80"

        delete_response = client.delete(
            f"/projects/{project_id}/tasks/1001/role-assignments/{assignment_id}",
            headers=headers,
        )
        assert delete_response.status_code == 204


def test_task_role_assignment_defaults_to_project_root_cost_code() -> None:
    """Issue #63 (E6-02): a role assignment created without an explicit
    `cost_code_id` is attached to the project's active root cost code by default."""
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))
        labor_role_id, _ = _seed_roles()
        root_id = _root_cost_code_id(client, headers, project_id)

        create_response = client.post(
            f"/projects/{project_id}/tasks/1001/role-assignments",
            json={"role_id": labor_role_id, "quantity": "1", "hours": "1"},
            headers=headers,
        )
        assert create_response.status_code == 201
        assert create_response.json()["cost_code_id"] == root_id


def test_task_role_assignment_accepts_explicit_sub_cost_code_and_rejects_foreign_one() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))
        labor_role_id, _ = _seed_roles()
        root_id = _root_cost_code_id(client, headers, project_id)

        sub_code_response = client.post(
            f"/projects/{project_id}/cost-codes",
            json={"code": "LOT-MO", "name": "Lot main d'oeuvre", "parent_id": root_id},
            headers=headers,
        )
        assert sub_code_response.status_code == 201
        sub_code_id = cast(int, sub_code_response.json()["id"])

        create_response = client.post(
            f"/projects/{project_id}/tasks/1001/role-assignments",
            json={
                "role_id": labor_role_id,
                "quantity": "1",
                "hours": "1",
                "cost_code_id": sub_code_id,
            },
            headers=headers,
        )
        assert create_response.status_code == 201
        assignment_id = cast(int, create_response.json()["id"])
        assert create_response.json()["cost_code_id"] == sub_code_id

        other_project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))
        other_root_id = _root_cost_code_id(client, headers, other_project_id)

        rejected_create_response = client.post(
            f"/projects/{project_id}/tasks/1002/role-assignments",
            json={
                "role_id": labor_role_id,
                "quantity": "1",
                "hours": "1",
                "cost_code_id": other_root_id,
            },
            headers=headers,
        )
        assert rejected_create_response.status_code == 400

        rejected_update_response = client.patch(
            f"/projects/{project_id}/tasks/1001/role-assignments/{assignment_id}",
            json={"cost_code_id": other_root_id},
            headers=headers,
        )
        assert rejected_update_response.status_code == 400

        accepted_update_response = client.patch(
            f"/projects/{project_id}/tasks/1001/role-assignments/{assignment_id}",
            json={"cost_code_id": root_id},
            headers=headers,
        )
        assert accepted_update_response.status_code == 200
        assert accepted_update_response.json()["cost_code_id"] == root_id


def test_task_role_assignment_explicit_null_cost_code_id_resolves_to_root_not_cleared() -> None:
    """Review finding on #63: `{"cost_code_id": null}` must resolve back to the
    project's root, exactly like an omitted field at create time -- never leave the
    assignment with no attachment at all."""
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))
        labor_role_id, _ = _seed_roles()
        root_id = _root_cost_code_id(client, headers, project_id)

        sub_code_response = client.post(
            f"/projects/{project_id}/cost-codes",
            json={"code": "LOT-MO-NULL", "name": "Lot", "parent_id": root_id},
            headers=headers,
        )
        sub_code_id = cast(int, sub_code_response.json()["id"])

        create_response = client.post(
            f"/projects/{project_id}/tasks/1001/role-assignments",
            json={
                "role_id": labor_role_id,
                "quantity": "1",
                "hours": "1",
                "cost_code_id": sub_code_id,
            },
            headers=headers,
        )
        assignment_id = cast(int, create_response.json()["id"])
        assert create_response.json()["cost_code_id"] == sub_code_id

        null_update_response = client.patch(
            f"/projects/{project_id}/tasks/1001/role-assignments/{assignment_id}",
            json={"cost_code_id": None},
            headers=headers,
        )
        assert null_update_response.status_code == 200
        assert null_update_response.json()["cost_code_id"] == root_id


def test_task_role_assignment_rejects_a_deactivated_cost_code() -> None:
    """Review finding on #63: a deactivated cost code must never accept a new
    attachment, even though it may still be referenced by pre-existing assignments."""
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))
        labor_role_id, _ = _seed_roles()
        root_id = _root_cost_code_id(client, headers, project_id)

        sub_code_response = client.post(
            f"/projects/{project_id}/cost-codes",
            json={"code": "LOT-MO-INACTIVE", "name": "Lot", "parent_id": root_id},
            headers=headers,
        )
        sub_code_id = cast(int, sub_code_response.json()["id"])

        deactivate_response = client.delete(
            f"/projects/{project_id}/cost-codes/{sub_code_id}", headers=headers
        )
        assert deactivate_response.status_code == 204

        rejected_create_response = client.post(
            f"/projects/{project_id}/tasks/1001/role-assignments",
            json={
                "role_id": labor_role_id,
                "quantity": "1",
                "hours": "1",
                "cost_code_id": sub_code_id,
            },
            headers=headers,
        )
        assert rejected_create_response.status_code == 400


def test_task_role_assignment_reports_a_clean_error_if_project_has_no_active_root() -> None:
    """Review finding on #63: the #62/E6-01 invariant (every project always has an
    active root cost code) is guaranteed by application logic, not a DB constraint --
    if it were ever violated, resolve_cost_code_id must surface a clean, documented
    error rather than a bare, unhandled NoResultFound. Deactivating the root directly
    at the ORM layer (bypassing the API, which always refuses this) is the only way
    to construct that state for this test."""
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))
        labor_role_id, _ = _seed_roles()
        root_id = _root_cost_code_id(client, headers, project_id)

        session_factory = get_session_factory()
        with session_factory() as session:
            root = session.query(ProjectCostCode).filter(ProjectCostCode.id == root_id).one()
            root.is_active = False
            session.add(root)
            session.commit()

        response = client.post(
            f"/projects/{project_id}/tasks/1001/role-assignments",
            json={"role_id": labor_role_id, "quantity": "1", "hours": "1"},
            headers=headers,
        )
        assert response.status_code == 500


def test_role_filter_can_include_descendant_nodes() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        labor_role_id, _ = _seed_roles()

        session_factory = get_session_factory()
        with session_factory() as session:
            role = session.query(ResourceRole).filter(ResourceRole.id == labor_role_id).one()
            root_id = session.query(ResourceNode).filter(ResourceNode.code == "DIRECTION").one().id

        direct_response = client.get(f"/resources/roles?node_id={root_id}", headers=headers)
        descendants_response = client.get(
            f"/resources/roles?node_id={root_id}&include_descendants=true",
            headers=headers,
        )
        assert direct_response.status_code == 200
        assert descendants_response.status_code == 200
        assert direct_response.json()["items"] == []
        assert role.id in [item["id"] for item in descendants_response.json()["items"]]


def _draft_structure_payload() -> dict[str, Any]:
    return {
        "posts": [
            {
                "key": "design",
                "name": "Design",
                "lots": [
                    {
                        "key": "specification",
                        "name": "Specification",
                        "deliverables": [{"key": "requirements", "name": "Requirements"}],
                    }
                ],
            }
        ]
    }


def test_planning_structure_draft_save_and_read_round_trip() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        create = client.post("/projects", json={"name": "Draft round trip"}, headers=headers)
        assert create.status_code == 201
        project_id = cast(int, create.json()["id"])
        draft_path = f"/projects/{project_id}/planning-structure/draft"

        payload = _draft_structure_payload()
        saved = client.put(draft_path, json=payload, headers=headers)
        assert saved.status_code == 200

        read = client.get(draft_path, headers=headers)
        assert read.status_code == 200
        body = cast(dict[str, Any], read.json())
        assert body["planning_id"] == saved.json()["planning_id"]
        assert body["structure"] == payload


def test_planning_structure_draft_read_returns_404_when_absent() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        create = client.post("/projects", json={"name": "No draft"}, headers=headers)
        assert create.status_code == 201
        project_id = cast(int, create.json()["id"])

        read = client.get(f"/projects/{project_id}/planning-structure/draft", headers=headers)
        assert read.status_code == 404


def test_planning_structure_draft_read_returns_409_for_invalid_json() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        create = client.post("/projects", json={"name": "Invalid draft json"}, headers=headers)
        assert create.status_code == 201
        project_id = cast(int, create.json()["id"])

        with get_session_factory()() as session:
            planning = WfPlanning(
                project_id=project_id,
                version_number=1,
                status="draft",
                structure_draft_json="{invalid-json",
            )
            session.add(planning)
            session.commit()

        read = client.get(f"/projects/{project_id}/planning-structure/draft", headers=headers)
        assert read.status_code == 409
        assert read.json()["detail"] == {"code": "GENERIC_ERROR"}


def test_planning_structure_draft_read_returns_409_for_invalid_schema() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        create = client.post("/projects", json={"name": "Invalid draft schema"}, headers=headers)
        assert create.status_code == 201
        project_id = cast(int, create.json()["id"])

        with get_session_factory()() as session:
            planning = WfPlanning(
                project_id=project_id,
                version_number=1,
                status="draft",
                structure_draft_json='{"posts":"not-a-list"}',
            )
            session.add(planning)
            session.commit()

        read = client.get(f"/projects/{project_id}/planning-structure/draft", headers=headers)
        assert read.status_code == 409
        assert read.json()["detail"] == {"code": "GENERIC_ERROR"}


def test_delete_project_clears_reference_estimate_before_deleting_estimates() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        project_id, _ = _seed_projects_and_tasks(_current_user_id(client, headers))
        create_response = client.post(
            f"/projects/{project_id}/estimates",
            json={"kind": "initial", "currency_code": "EUR"},
            headers=headers,
        )
        assert create_response.status_code == 201
        estimate_id = cast(int, create_response.json()["id"])

        # Point the project at the estimate so deletion must clear the FK first.
        with get_session_factory()() as session:
            project = session.get(MsProject, project_id)
            assert project is not None
            project.reference_estimate_id = estimate_id
            session.commit()

        deleted = client.delete(f"/projects/{project_id}", headers=headers)
        assert deleted.status_code == 204

        with get_session_factory()() as session:
            assert session.get(MsProject, project_id) is None
            assert session.query(Estimate).filter(Estimate.id == estimate_id).count() == 0


def test_generate_planning_structure_syncs_persisted_draft() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        create = client.post("/projects", json={"name": "Draft sync"}, headers=headers)
        assert create.status_code == 201
        project_id = cast(int, create.json()["id"])
        draft_path = f"/projects/{project_id}/planning-structure/draft"

        structure_a = _draft_structure_payload()
        assert client.put(draft_path, json=structure_a, headers=headers).status_code == 200

        structure_b: dict[str, Any] = {
            "posts": [
                {
                    "key": "build",
                    "name": "Build",
                    "lots": [
                        {
                            "key": "assembly",
                            "name": "Assembly",
                            "deliverables": [{"key": "module", "name": "Module"}],
                        }
                    ],
                }
            ]
        }
        generated = client.post(
            f"/projects/{project_id}/planning-structure",
            json=structure_b,
            headers=headers,
        )
        assert generated.status_code == 201

        read = client.get(draft_path, headers=headers)
        assert read.status_code == 200
        assert read.json()["structure"] == structure_b


def test_planning_tasks_are_returned_depth_first() -> None:
    with TestClient(app) as client:
        headers = _auth_headers(client)
        create = client.post("/projects", json={"name": "Depth first"}, headers=headers)
        assert create.status_code == 201
        project_id = cast(int, create.json()["id"])

        structure: dict[str, Any] = {
            "posts": [
                {
                    "key": "p1",
                    "name": "Post 1",
                    "lots": [
                        {
                            "key": "l1",
                            "name": "Lot 1",
                            "deliverables": [
                                {"key": "d1", "name": "Deliverable 1"},
                                {"key": "d2", "name": "Deliverable 2"},
                            ],
                        },
                        {
                            "key": "l2",
                            "name": "Lot 2",
                            "deliverables": [{"key": "d3", "name": "Deliverable 3"}],
                        },
                    ],
                },
                {
                    "key": "p2",
                    "name": "Post 2",
                    "lots": [
                        {
                            "key": "l3",
                            "name": "Lot 3",
                            "deliverables": [{"key": "d4", "name": "Deliverable 4"}],
                        }
                    ],
                },
            ]
        }
        generated = client.post(
            f"/projects/{project_id}/planning-structure",
            json=structure,
            headers=headers,
        )
        assert generated.status_code == 201

        tasks = cast(
            list[dict[str, Any]],
            client.get(f"/projects/{project_id}/tasks", headers=headers).json()["items"],
        )
        outlines = [task["outline_number"] for task in tasks]
        # Parent immediately followed by its children, in local position order.
        assert outlines[:4] == ["1", "1.1", "1.1.1", "1.1.2"]
        # A second deliverable of lot 1 precedes lot 2 and the next post subtree.
        assert outlines.index("1.1.2") < outlines.index("1.2")
        assert outlines.index("1.2.1") < outlines.index("2")


def _seed_complete_setup() -> None:
    """Seed every prerequisite GET /projects/setup-warnings checks for: an
    active default calendar with a working day, an active cost category, and
    an active resource role."""
    session_factory = get_session_factory()
    with session_factory() as session:
        calendar = Calendar(code="STANDARD", name="Standard", weeks_per_year=47, is_default=True)
        session.add(calendar)
        session.flush()
        session.add_all(
            CalendarWeekday(
                calendar_id=calendar.id,
                day_type=day_type,
                hours_per_day=Decimal("0.00") if day_type in (1, 7) else Decimal("7.00"),
            )
            for day_type in range(1, 8)
        )

        node = ResourceNode(code="IT", name="Departement informatique")
        session.add(node)
        session.flush()
        cost_type = CostType(code="MO", name="Main d'oeuvre", kind="labor")
        session.add(cost_type)
        session.flush()
        category = CostCategory(
            cost_type_id=cost_type.id,
            accounting_code="DEV",
            category_code="IDEX",
            name="Developpement",
        )
        session.add(category)
        session.flush()
        role = ResourceRole(node_id=node.id, cost_category_id=category.id, name="Developpeur")
        session.add(role)
        session.commit()


class TestProjectSetupWarnings:
    """GET /projects/setup-warnings (issue #109): non-blocking global
    setup-prerequisite checks surfaced before project creation."""

    def test_reports_no_warnings_when_setup_is_complete(self) -> None:
        with TestClient(app) as client:
            headers = _auth_headers(client)
            _seed_complete_setup()

            response = client.get("/projects/setup-warnings", headers=headers)

            assert response.status_code == 200
            assert response.json() == {"warnings": []}

    def test_reports_missing_default_calendar(self) -> None:
        with TestClient(app) as client:
            headers = _auth_headers(client)
            _seed_complete_setup()
            with get_session_factory()() as session:
                session.query(Calendar).update({Calendar.is_default: False})
                session.commit()

            response = client.get("/projects/setup-warnings", headers=headers)

            assert response.status_code == 200
            codes = [warning["code"] for warning in response.json()["warnings"]]
            assert codes == ["no_default_calendar"]

    def test_reports_default_calendar_without_working_day(self) -> None:
        with TestClient(app) as client:
            headers = _auth_headers(client)
            _seed_complete_setup()
            with get_session_factory()() as session:
                session.query(CalendarWeekday).update(
                    {CalendarWeekday.hours_per_day: Decimal("0.00")}
                )
                session.commit()

            response = client.get("/projects/setup-warnings", headers=headers)

            assert response.status_code == 200
            codes = [warning["code"] for warning in response.json()["warnings"]]
            assert codes == ["default_calendar_has_no_working_day"]

    def test_reports_no_active_cost_category(self) -> None:
        with TestClient(app) as client:
            headers = _auth_headers(client)
            _seed_complete_setup()
            with get_session_factory()() as session:
                session.query(CostCategory).update({CostCategory.is_active: False})
                session.commit()

            response = client.get("/projects/setup-warnings", headers=headers)

            assert response.status_code == 200
            codes = [warning["code"] for warning in response.json()["warnings"]]
            assert codes == ["no_active_cost_category"]

    def test_reports_no_active_resource_role(self) -> None:
        with TestClient(app) as client:
            headers = _auth_headers(client)
            _seed_complete_setup()
            with get_session_factory()() as session:
                session.query(ResourceRole).update({ResourceRole.is_active: False})
                session.commit()

            response = client.get("/projects/setup-warnings", headers=headers)

            assert response.status_code == 200
            codes = [warning["code"] for warning in response.json()["warnings"]]
            assert codes == ["no_active_resource_role"]

    def test_reports_every_missing_prerequisite_together(self) -> None:
        with TestClient(app) as client:
            headers = _auth_headers(client)
            # Nothing seeded at all: every one of the 3 checks should fire.

            response = client.get("/projects/setup-warnings", headers=headers)

            assert response.status_code == 200
            codes = {warning["code"] for warning in response.json()["warnings"]}
            assert codes == {
                "no_default_calendar",
                "no_active_cost_category",
                "no_active_resource_role",
            }

    def test_does_not_block_project_creation(self) -> None:
        """Issue #109: the setup-warnings check is purely advisory --
        POST /projects must keep succeeding regardless of what it reports."""
        with TestClient(app) as client:
            headers = _auth_headers(client)

            warnings = client.get("/projects/setup-warnings", headers=headers)
            assert warnings.status_code == 200
            assert warnings.json()["warnings"]

            created = client.post(
                "/projects", json={"name": "Created despite warnings"}, headers=headers
            )
            assert created.status_code == 201
